"""
DeltaBalance — migration/migrar_transacciones_noe.py

Importa las transacciones de las hojas mensuales de la planilla de Noe
(migration/sources/Finanzas.xlsx, hojas 2024-02 a 2026-09) a
data/deltabalanceNR.db. Misma mecánica que migration/migrar_transacciones.py
(el de Bruno): SQL directo vía DatabaseManager, sin pasar por ningún
service. Lo que cambia:

- Cuentas: CUENTAS_DEBITO. Las que falten se crean con su cuentas_saldos en
  ARS; CAJA EFECTIVO ya la trae el seed ("Caja Efectivo") y se reusa.
- Saldos iniciales (SALDOS_INICIALES): saldo_inicial_minor de esa cuenta en
  esa moneda. cuentas_saldos no tiene fecha: el saldo inicial es el de
  antes de la primera transacción, y la primera hoja es 2024-02, así que es
  el saldo al 2024-02-01. Si la fila ya tiene otro saldo distinto de 0 (ej.
  corregido a mano en la app), no se pisa: se avisa.
- Subcategoría → (categoría, tag) (MAPA_SUBCATEGORIA): varias subcategorías
  de la planilla van a una misma categoría con una etiqueta en
  transacciones.tag (ej. UBER → INGRESO VARIABLE con tag COBRANZAS).
- Bancos salteados: OTRO, AHORRO (error de carga en la planilla) y vacío.
- Categorías que el seed no trae y MAPA_SUBCATEGORIA necesita
  (CATEGORIAS_NUEVAS): con la misma categoría principal y tipo que les da
  el script de Bruno.
- Notas: "MIGRADO DESDE EXCEL NOE — HOJA YYYY-MM, FILA N".

Igual que en el de Bruno — sus funciones se importan, no se copian, así
que cualquier ajuste allá vale acá: detección de columnas por encabezado
(_encontrar_header()), corrección de fechas (_resolver_fecha(): si la
fecha no cae en el mes de la hoja, se corrige el año), monedas (columna
MONEDA si está, CHL → CLP, decimales de cada moneda), filas resumen
ignoradas, tipo_movimiento por el signo del monto, e idempotencia contando
apariciones: concepto + fecha + cuenta_id + monto_minor + moneda_id (+
tipo_movimiento). El porqué de cada regla está en el docstring de
migrar_transacciones.py.

--- Transacción y seguridad ---

Por default es DRY-RUN: todo corre dentro de una transacción (cada hoja en
un SAVEPOINT), se imprime el reporte completo y al final se hace rollback.
Con --confirmar: backup de la DB, commit del setup (categorías, cuentas,
saldos iniciales) y commit de cada hoja al terminarla; si una hoja falla,
rollback de esa hoja sola y se sigue con la siguiente. Nunca corre contra
data/deltabalance.db ni contra la base de Bruno (data/deltabalanceBZ.db),
aunque se pase por --db-path. DatabaseManager.inicializar() crea la DB si
no existe (schema + seed), igual que al abrir la app.

Requiere openpyxl >= 3.1 (pip install openpyxl).

Uso:
    # 1) Dry-run (no escribe nada, solo reporta):
    python migration/migrar_transacciones_noe.py

    # 2) Aplicar de verdad (hace backup antes):
    python migration/migrar_transacciones_noe.py --confirmar

    # Opcional: contra otra DB (ej. una copia de prueba)
    python migration/migrar_transacciones_noe.py --db-path /tmp/prueba_noe.db --confirmar

Este script NO lo ejecuta Claude Code (CLAUDE.md §0.1) — lo corre el
usuario a mano.
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
from pathlib import Path
from typing import Any, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    import openpyxl
except ImportError:
    sys.exit("❌ Falta openpyxl — instalalo en el entorno del proyecto:  pip install openpyxl")

from db.database import DatabaseManager, to_minor
from migration.migrar_transacciones import (
    COL_BANCO,
    COL_CONCEPTO,
    COL_FECHA,
    COL_MONEDA,
    COL_MONTO,
    COL_SUBCATEGORIA,
    MONEDA_ALIAS,
    MONEDA_DEFAULT,
    PATRON_HOJA_MENSUAL,
    MigracionError,
    _categorias_activas,
    _clave,
    _contadores_vacios,
    _encontrar_header,
    _mapa_categorias,
    _parsear_monto,
    _resolver_fecha,
    _texto,
)

RAIZ = Path(__file__).resolve().parent.parent
DB_PATH = RAIZ / "data" / "deltabalanceNR.db"
# La DB real de la app y la de Bruno: este script nunca las toca.
DBS_PROHIBIDAS = (RAIZ / "data" / "deltabalance.db", RAIZ / "data" / "deltabalanceBZ.db")
SOURCES_DIR = Path(__file__).resolve().parent / "sources"
XLSX_PATH = SOURCES_DIR / "Finanzas.xlsx"
# Hojas a procesar (inclusive): 32 meses.
HOJA_DESDE = "2024-02"
HOJA_HASTA = "2026-09"
NOTA_MIGRACION = "MIGRADO DESDE EXCEL NOE — HOJA {hoja}, FILA {fila}"

# --- Saldos iniciales al 2024-02-01: (cuenta, moneda, saldo en minor units) ---
SALDOS_INICIALES = [
    ("MP",            "ARS", 1560487),   # $15,604.87
    ("CAJA EFECTIVO", "ARS",    5000),   # $50.00
]

# --- Cuentas (las que falten se crean) ---
CUENTAS_DEBITO = [
    "MP", "NACION", "NARANJA X", "BILLETERA ER", "CAJA EFECTIVO", "BERSA",
]
# Tipo con el que se crea una cuenta que todavía no existe (mismo criterio
# que el script de Bruno: CAJA EFECTIVO es 'efectivo', como la del seed).
TIPO_CUENTA_NUEVA = {"CAJA EFECTIVO": "efectivo"}
TIPO_CUENTA_DEFAULT = "debito"

# --- Mapeo de banco → cuenta ---
MAPA_BANCO = {
    "MP":                   "MP",
    "NACION":               "NACION",
    "NARANJA X":            "NARANJA X",
    "BILLETERA ER":         "BILLETERA ER",
    "BILLETERA ENTRE RIOS": "BILLETERA ER",
    "CAJA":                 "CAJA EFECTIVO",
    "BERSA":                "BERSA",
}
# OTRO: fuera de las cuentas; AHORRO: error de carga en la planilla; "" = BANCO vacío.
BANCOS_IGNORAR = {"OTRO", "AHORRO", ""}

# --- Mapeo de subcategoría → (categoría DB, tag) (None = ignorar la fila) ---
MAPA_SUBCATEGORIA: dict[str, Optional[tuple[str, Optional[str]]]] = {
    # Con tag
    "UBER":                ("INGRESO VARIABLE",    "COBRANZAS"),
    "DIDI":                ("INGRESO VARIABLE",    "COBRANZAS"),
    "VELAS":               ("INGRESO VARIABLE",    "COBRANZAS"),
    "PACIFIC SHOES":       ("INGRESO VARIABLE",    "COBRANZAS"),
    "COBRANZAS":           ("INGRESO VARIABLE",    "COBRANZAS"),
    "COBRANZAS MP":        ("INGRESO VARIABLE",    "COBRANZAS"),
    "COBRANZAS NX/MP":     ("INGRESO VARIABLE",    "COBRANZAS"),
    "CUMPLE MAMI":         ("INGRESO VARIABLE",    "CUMPLE MAMI"),
    "DEUDA BRUNO":         ("DEUDA",               "BRUNO"),
    "DEUDA MAMI":          ("DEUDA",               "MAMI"),
    # Sin tag
    "PICHOS":              ("MASCOTAS",            None),
    "AHORRO AUTO":         ("AHORRO/INVERSIÓN",    None),
    "COMPARTIDO":          ("EGRESO VARIABLE",     None),
    "TARJETA":             ("SERVICIOS",           None),
    "INGRESO FIJO":        ("SUELDO",              None),
    "INGRESO VAR":         ("INGRESO VARIABLE",    None),
    "INGRESO VARIABLE":    ("INGRESO VARIABLE",    None),
    "EGRESO FIJO":         ("SERVICIOS",           None),
    "EGRESO VARIABLE":     ("EGRESO VARIABLE",     None),
    "AUTOTRANSFERENCIA":   ("AUTOTRANSFERENCIA",   None),
    "TRANSFERENCIA":       ("AUTOTRANSFERENCIA",   None),
    "TRANSFERENCIA OTRO":  ("EGRESO VARIABLE",     None),
    "HOGAR":               ("HOGAR",               None),
    "OCIO":                ("OCIO",                None),
    "PERSONAL":            ("BIENESTAR Y DEPORTE", None),
    "REGALOS":             ("REGALOS",             None),
    "PROVEDURIA":          ("ALIMENTOS",           None),
    "TRANSPORTE":          ("TRANSPORTE / AUTO",   None),
    "RENDIMIENTOS":        ("RENDIMIENTOS",        None),
    "INVERSIONES":         ("INVERSIONES",         None),
    "REINTEGRO PROMOCION": ("REINTEGRO PROMOCION", None),
    "VACACIONES":          ("VACACIONES",          None),
    "RESERVA":             None,  # ignorar
    "DEUDA A COBRAR": ("DEUDA", None),
    "DEUDA A PAGAR":  ("DEUDA", None),
}

# --- Filas a ignorar: resúmenes del sistema anterior, no transacciones reales ---
CONCEPTOS_IGNORAR = {
    "INGRESO FIJO", "EGRESO FIJO", "EGRESO VARIABLE", "INGRESO VARIABLE",
    "CREDITO ESTE MES", "CREDITO PROXIMO MES", "DEUDAS A COBRAR",
    "DEUDAS A PAGAR", "INGRESOS", "EGRESOS", "SALDO INICIAL", "DEUDAS",
    "INGRESO VAR", "INGRESO", "EGRESO",
}

# --- Categorías a crear si no existen: las destino de MAPA_SUBCATEGORIA que
# el seed no trae, con la principal y el tipo del script de Bruno ---
CATEGORIAS_NUEVAS = [
    ("EGRESOS VARIABLES", "EGRESO VARIABLE",     "egreso"),
    ("EGRESOS VARIABLES", "ALIMENTOS",           "egreso"),
    ("EGRESOS VARIABLES", "REGALOS",             "egreso"),
    ("EGRESOS VARIABLES", "MASCOTAS",            "egreso"),
    ("EGRESOS VARIABLES", "VACACIONES",          "egreso"),
    ("EGRESOS VARIABLES", "REINTEGRO PROMOCION", "egreso"),
    ("INGRESOS",          "INGRESO VARIABLE",    "ingreso"),
]


# ============================================================
# CHEQUEOS DE CONFIGURACIÓN (antes de tocar la DB)
# ============================================================

def _validar_configuracion() -> None:
    """Toda cuenta de MAPA_BANCO y de SALDOS_INICIALES tiene que estar en CUENTAS_DEBITO."""
    cuentas = {_clave(c) for c in CUENTAS_DEBITO}
    sueltas = sorted(
        {destino for destino in MAPA_BANCO.values() if _clave(destino) not in cuentas}
        | {cuenta for cuenta, _, _ in SALDOS_INICIALES if _clave(cuenta) not in cuentas}
    )
    if sueltas:
        raise MigracionError(f"Cuentas usadas que no están en CUENTAS_DEBITO: {', '.join(sueltas)}.")


# ============================================================
# SETUP: CATEGORÍAS, CUENTAS Y SALDOS INICIALES
# ============================================================

def _crear_categorias(conn: sqlite3.Connection) -> None:
    print("\n[CATEGORÍAS] Verificando categorías...")
    for principal, sub, tipo in CATEGORIAS_NUEVAS:
        existente = next(
            (c for c in _categorias_activas(conn) if _clave(c["subcategoria"]) == _clave(sub)), None,
        )
        if existente is not None:
            ubicacion = f"{existente['categoria_principal']} · {existente['subcategoria']}"
            como = "" if ubicacion == f"{principal} · {sub}" else f" (como '{ubicacion}' — se reusa, no se duplica)"
            print(f"  ⏭️  {sub} ya existe{como}")
            continue
        cur = conn.execute(
            "INSERT OR IGNORE INTO categorias (categoria_principal, subcategoria, tipo) VALUES (?, ?, ?);",
            (principal, sub, tipo),
        )
        if cur.rowcount == 0:
            # Existe con ese nombre exacto pero desactivada: no se revierte un
            # soft-delete hecho a propósito — sus filas quedan como error.
            print(f"  ⚠️  {principal} · {sub} existe DESACTIVADA — no se reactiva")
        else:
            print(f"  ✅ {sub} creada")


def _asegurar_cuentas(conn: sqlite3.Connection, moneda_ars_id: int) -> dict[str, str]:
    """Nombre de CUENTAS_DEBITO → id. Crea las que falten; las que existen (por nombre, sin tildes ni mayúsculas) no se tocan."""
    print("\n[CUENTAS] Verificando cuentas...")
    ids: dict[str, str] = {}
    for nombre in CUENTAS_DEBITO:
        existentes = {_clave(r["nombre"]): r for r in conn.execute("SELECT id, nombre FROM cuentas;")}
        existente = existentes.get(_clave(nombre))
        if existente is not None:
            ids[nombre] = existente["id"]
            print(f"  ⏭️  {nombre} ya existe (como '{existente['nombre']}')")
            continue
        tipo = TIPO_CUENTA_NUEVA.get(nombre, TIPO_CUENTA_DEFAULT)
        conn.execute("INSERT OR IGNORE INTO cuentas (nombre, tipo) VALUES (?, ?);", (nombre, tipo))
        cuenta_id = conn.execute("SELECT id FROM cuentas WHERE nombre = ?;", (nombre,)).fetchone()["id"]
        conn.execute(
            "INSERT OR IGNORE INTO cuentas_saldos (cuenta_id, moneda_id, saldo_inicial_minor) VALUES (?, ?, 0);",
            (cuenta_id, moneda_ars_id),
        )
        ids[nombre] = cuenta_id
        print(f"  ✅ {nombre} creada ({tipo})")
    print(f"  {len(ids)} cuentas listas.")
    return ids


def _texto_monto(minor: int, moneda: sqlite3.Row) -> str:
    return f"{moneda['codigo']} ${minor / 10 ** moneda['decimales']:,.{moneda['decimales']}f}"


def _setear_saldos_iniciales(
    conn: sqlite3.Connection, ids_cuentas: dict[str, str], monedas: dict[str, sqlite3.Row],
) -> None:
    """SALDOS_INICIALES → cuentas_saldos.saldo_inicial_minor (ver docstring del módulo: un saldo distinto de 0 no se pisa)."""
    print(f"\n[SALDOS INICIALES] Al {HOJA_DESDE}-01...")
    for nombre, codigo, saldo_minor in SALDOS_INICIALES:
        moneda = monedas.get(_clave(codigo))
        if moneda is None:
            raise MigracionError(f"La moneda {codigo} del saldo inicial de {nombre} no existe en la DB.")
        cuenta_id = ids_cuentas[nombre]
        fila = conn.execute(
            "SELECT saldo_inicial_minor FROM cuentas_saldos WHERE cuenta_id = ? AND moneda_id = ?;",
            (cuenta_id, moneda["id"]),
        ).fetchone()
        objetivo = _texto_monto(saldo_minor, moneda)
        if fila is None:
            conn.execute(
                "INSERT INTO cuentas_saldos (cuenta_id, moneda_id, saldo_inicial_minor) VALUES (?, ?, ?);",
                (cuenta_id, moneda["id"], saldo_minor),
            )
            print(f"  ✅ {nombre}: {objetivo} (la cuenta ahora opera en {moneda['codigo']})")
        elif fila["saldo_inicial_minor"] == saldo_minor:
            print(f"  ⏭️  {nombre}: ya tiene {objetivo}")
        elif fila["saldo_inicial_minor"] == 0:
            conn.execute(
                "UPDATE cuentas_saldos SET saldo_inicial_minor = ? WHERE cuenta_id = ? AND moneda_id = ?;",
                (saldo_minor, cuenta_id, moneda["id"]),
            )
            print(f"  ✅ {nombre}: {_texto_monto(0, moneda)} → {objetivo}")
        else:
            print(
                f"  ⚠️  {nombre}: ya tiene {_texto_monto(fila['saldo_inicial_minor'], moneda)} "
                f"(no {objetivo}) — NO se pisa; si hay que cambiarlo, desde la app"
            )


# ============================================================
# HOJAS MENSUALES
# ============================================================

def _procesar_hoja(
    conn: sqlite3.Connection,
    ws,
    hoja: str,
    ids_cuentas: dict[str, str],
    categorias: dict[str, sqlite3.Row],
    monedas: dict[str, sqlite3.Row],
    ocurrencias: dict[tuple, int],
) -> dict[str, int]:
    """
    Procesa una hoja completa (el caller maneja su transacción). Mismo
    recorrido que migrar_transacciones._procesar_hoja(), con el tag de
    MAPA_SUBCATEGORIA. `ocurrencias` acumula, para todo el run, cuántas
    veces apareció cada clave de idempotencia en la planilla.
    """
    cont = _contadores_vacios()
    anio_hoja, mes_hoja = (int(x) for x in PATRON_HOJA_MENSUAL.fullmatch(hoja).groups())
    mapa_banco = {_clave(k): v for k, v in MAPA_BANCO.items()}
    mapa_subcategoria = {_clave(k): v for k, v in MAPA_SUBCATEGORIA.items()}
    ignorar_concepto = {_clave(c) for c in CONCEPTOS_IGNORAR}
    ignorar_banco = {_clave(b) for b in BANCOS_IGNORAR}

    fila_header, columnas = _encontrar_header(ws, hoja)
    tiene_moneda = COL_MONEDA in columnas
    if not tiene_moneda and anio_hoja >= 2025:
        print(f"  ⚠️  La hoja {hoja} no tiene columna {COL_MONEDA} — todo en {MONEDA_DEFAULT}")

    for fila_excel, valores in enumerate(
        ws.iter_rows(min_row=fila_header + 1, values_only=True), start=fila_header + 1,
    ):
        def celda(col: str) -> Any:
            indice = columnas.get(col)
            return valores[indice] if indice is not None and indice < len(valores) else None

        concepto = _texto(celda(COL_CONCEPTO))
        if not concepto:
            continue  # fila vacía
        cont["total"] += 1
        etiqueta = f"[{fila_excel:03d}] {concepto}"

        if _clave(concepto) in ignorar_concepto:
            cont["ignoradas"] += 1
            continue

        # --- Banco ---
        banco = _texto(celda(COL_BANCO))
        if _clave(banco) in ignorar_banco:
            cont["saltadas_banco_subcat"] += 1
            continue
        nombre_cuenta = mapa_banco.get(_clave(banco))
        if nombre_cuenta is None:
            cont["saltadas_banco_subcat"] += 1
            print(f'  ⚠️  {etiqueta} | BANCO DESCONOCIDO: "{banco}" — salteando')
            continue
        cuenta_id = ids_cuentas[nombre_cuenta]

        # --- Subcategoría → categoría + tag ---
        subcategoria = _texto(celda(COL_SUBCATEGORIA))
        if not subcategoria:
            cont["saltadas_banco_subcat"] += 1
            continue
        clave_sub = _clave(subcategoria)
        if clave_sub not in mapa_subcategoria:
            cont["saltadas_banco_subcat"] += 1
            print(f'  ⚠️  {etiqueta} | SUBCATEGORÍA DESCONOCIDA: "{subcategoria}" — salteando')
            continue
        destino = mapa_subcategoria[clave_sub]
        if destino is None:  # RESERVA
            cont["saltadas_banco_subcat"] += 1
            continue
        nombre_categoria, tag = destino
        categoria = categorias.get(_clave(nombre_categoria))
        if categoria is None:
            cont["errores"] += 1
            print(f"  ❌ {etiqueta} | LA CATEGORÍA DESTINO '{nombre_categoria}' NO EXISTE (ACTIVA) EN LA DB — salteando")
            continue

        # --- Monto ---
        valor_monto = celda(COL_MONTO)
        monto = _parsear_monto(valor_monto)
        if monto is None or monto == 0:
            cont["saltadas_monto"] += 1
            if isinstance(valor_monto, str) and valor_monto.strip():
                print(f"  ⚠️  {etiqueta} | MONTO NO NUMÉRICO: {valor_monto!r} — salteando")
            continue

        avisos: list[str] = []

        # --- Fecha (manda la hoja: ver migrar_transacciones._resolver_fecha()) ---
        fecha, aviso_fecha = _resolver_fecha(celda(COL_FECHA), anio_hoja, mes_hoja)
        if aviso_fecha:
            cont["fechas_corregidas"] += 1
            avisos.append(aviso_fecha)

        # --- Moneda ---
        codigo_moneda = MONEDA_DEFAULT
        if tiene_moneda:
            moneda_txt = _texto(celda(COL_MONEDA))
            codigo_moneda = MONEDA_ALIAS.get(moneda_txt, moneda_txt)
            if _clave(codigo_moneda) not in monedas:
                avisos.append(f"MONEDA DESCONOCIDA: \"{moneda_txt or '(vacía)'}\" → {MONEDA_DEFAULT}")
                codigo_moneda = MONEDA_DEFAULT
        moneda = monedas[_clave(codigo_moneda)]

        tipo_movimiento = "ingreso" if monto > 0 else "egreso"
        monto_minor = abs(to_minor(monto, moneda["decimales"]))
        if monto_minor == 0:
            # Monto real pero menor a la unidad mínima de la moneda (ej. 0.3 CLP).
            cont["saltadas_monto"] += 1
            print(f"  ⚠️  {etiqueta} | MONTO {monto} REDONDEA A 0 EN {moneda['codigo']} — salteando")
            continue

        # --- Idempotencia (contando apariciones) ---
        clave_idem = (concepto, fecha, cuenta_id, monto_minor, moneda["id"], tipo_movimiento)
        ocurrencias[clave_idem] = ocurrencias.get(clave_idem, 0) + 1
        en_db = conn.execute(
            """
            SELECT COUNT(*) AS n FROM transacciones
            WHERE concepto = ? AND fecha = ? AND cuenta_id = ? AND monto_minor = ?
              AND moneda_id = ? AND tipo_movimiento = ?;
            """,
            clave_idem,
        ).fetchone()["n"]
        if en_db >= ocurrencias[clave_idem]:
            cont["ya_existian"] += 1
            print(f"  ⏭️  {etiqueta} | YA EXISTE")
            continue

        # --- La cuenta tiene que operar en la moneda de la transacción ---
        cur = conn.execute(
            "INSERT OR IGNORE INTO cuentas_saldos (cuenta_id, moneda_id, saldo_inicial_minor) VALUES (?, ?, 0);",
            (cuenta_id, moneda["id"]),
        )
        if cur.rowcount:
            avisos.append(f"{nombre_cuenta} AHORA OPERA TAMBIÉN EN {moneda['codigo']} (cuentas_saldos)")

        conn.execute(
            """
            INSERT INTO transacciones
                (fecha, concepto, cuenta_id, categoria_id, moneda_id, tipo_movimiento, monto_minor, tag, notas)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
            """,
            (
                fecha, concepto, cuenta_id, categoria["id"], moneda["id"], tipo_movimiento, monto_minor, tag,
                NOTA_MIGRACION.format(hoja=hoja, fila=fila_excel),
            ),
        )
        cont["importadas"] += 1
        for aviso in avisos:
            print(f"  ⚠️  {etiqueta} | {aviso}")
        print(
            f"  ✅ {etiqueta} | {nombre_cuenta} | {moneda['codigo']} ${abs(monto):,.2f} | "
            f"{tipo_movimiento} | {fecha}" + (f" | TAG {tag}" if tag else "")
        )
    return cont


# ============================================================
# MAIN
# ============================================================

def _mostrar_ruta(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(RAIZ))
    except ValueError:
        return str(path)


def _hojas_esperadas() -> list[str]:
    """'YYYY-MM' de HOJA_DESDE a HOJA_HASTA, inclusive."""
    anio, mes = (int(x) for x in HOJA_DESDE.split("-"))
    hojas: list[str] = []
    while f"{anio:04d}-{mes:02d}" <= HOJA_HASTA:
        hojas.append(f"{anio:04d}-{mes:02d}")
        anio, mes = (anio + 1, 1) if mes == 12 else (anio, mes + 1)
    return hojas


def main() -> None:
    parser = argparse.ArgumentParser(description="Migra las transacciones de las hojas mensuales de la planilla de Noe.")
    parser.add_argument("--db-path", default=None, help=f"Ruta a la DB (default: {_mostrar_ruta(DB_PATH)}).")
    parser.add_argument("--confirmar", action="store_true", help="Aplica los cambios. Sin esto, solo dry-run.")
    parser.add_argument("--desde", default=None, help=f"Procesar solo hojas desde este mes (ej: 2026-01; default {HOJA_DESDE}).")
    args = parser.parse_args()

    db_path = Path(args.db_path) if args.db_path else DB_PATH
    for prohibida in DBS_PROHIBIDAS:
        if db_path.resolve() == prohibida.resolve():
            sys.exit(f"❌ Este script nunca corre contra {_mostrar_ruta(prohibida)} (es de Noe: {_mostrar_ruta(DB_PATH)}).")

    if not XLSX_PATH.exists():
        sys.exit(f"❌ No se encontró la planilla de Noe: {_mostrar_ruta(XLSX_PATH)}.")

    try:
        _validar_configuracion()
    except MigracionError as err:
        sys.exit(f"❌ {err}")

    modo = "APLICANDO CAMBIOS" if args.confirmar else "DRY-RUN (nada se escribe)"
    print(f"[INICIO] Migración de transacciones mensuales de NOE — {modo}")
    print(f"[XLSX] {_mostrar_ruta(XLSX_PATH)}")
    print(f"[DB] Conectando a {_mostrar_ruta(db_path)}")

    db = DatabaseManager(db_path)
    db.inicializar()  # idempotente: schema + migraciones; seed solo si la DB es nueva
    conn = db.conn
    # inicializar() puede dejar abierta una transacción implícita: se cierra
    # antes de pasar a control manual (BEGIN/SAVEPOINT/COMMIT explícitos).
    conn.commit()
    conn.isolation_level = None

    if args.confirmar:
        # El backup copia solo el archivo .db — con WAL, lo que siga en el -wal no estaría sin este checkpoint.
        conn.execute("PRAGMA wal_checkpoint(TRUNCATE);")
        db.hacer_backup()

    conn.execute("BEGIN;")
    try:
        monedas = {_clave(m["codigo"]): m for m in conn.execute("SELECT id, codigo, decimales FROM monedas;")}
        if _clave(MONEDA_DEFAULT) not in monedas:
            raise MigracionError(f"La moneda default {MONEDA_DEFAULT} no existe en la DB.")
        _crear_categorias(conn)
        categorias = _mapa_categorias(conn)
        ids_cuentas = _asegurar_cuentas(conn, monedas[_clave(MONEDA_DEFAULT)]["id"])
        _setear_saldos_iniciales(conn, ids_cuentas, monedas)

        destinos = {d[0] for d in MAPA_SUBCATEGORIA.values() if d is not None}
        faltantes = sorted(d for d in destinos if _clave(d) not in categorias)
        if faltantes:
            print(f"\n  ⚠️  Categorías destino que NO existen activas en la DB (sus filas van a Errores): {', '.join(faltantes)}")
    except Exception as err:
        conn.execute("ROLLBACK;")
        db.desconectar()
        print(f"\n❌ ERROR en el setup — rollback completo, no se escribió nada: {err}")
        if isinstance(err, MigracionError):
            sys.exit(1)
        raise

    if args.confirmar:
        conn.execute("COMMIT;")
        conn.execute("BEGIN;")

    totales = _contadores_vacios()
    hojas_ok = 0
    hojas_con_error: list[str] = []
    ocurrencias: dict[tuple, int] = {}
    esperadas = _hojas_esperadas()

    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    try:
        en_planilla = {n for n in wb.sheetnames if PATRON_HOJA_MENSUAL.fullmatch(n)}
        hojas = [h for h in esperadas if h in en_planilla]
        print(f"\n{len(hojas)} de {len(esperadas)} hojas mensuales ({HOJA_DESDE} … {HOJA_HASTA})")
        faltan = [h for h in esperadas if h not in en_planilla]
        if faltan:
            print(f"  ⚠️  No están en la planilla: {', '.join(faltan)}")
        fuera = sorted(en_planilla - set(esperadas))
        if fuera:
            print(f"  ⏭️  Fuera del rango, no se procesan: {', '.join(fuera)}")
        for hoja in hojas:
            if args.desde and hoja < args.desde:
                continue
            print(f"\n[HOJA {hoja}]")
            ocurrencias_antes = dict(ocurrencias)
            conn.execute("SAVEPOINT hoja;")
            try:
                cont = _procesar_hoja(conn, wb[hoja], hoja, ids_cuentas, categorias, monedas, ocurrencias)
            except Exception as err:
                conn.execute("ROLLBACK TO hoja;")
                conn.execute("RELEASE hoja;")
                ocurrencias = ocurrencias_antes
                hojas_con_error.append(hoja)
                print(f"  ❌ ERROR en la hoja {hoja} — rollback de la hoja, se sigue con la siguiente: {err!r}")
                continue
            conn.execute("RELEASE hoja;")
            if args.confirmar:
                conn.execute("COMMIT;")
                conn.execute("BEGIN;")
            hojas_ok += 1
            for k, v in cont.items():
                totales[k] += v
            print(
                f"  → {cont['total']} filas: {cont['importadas']} importadas, {cont['ya_existian']} ya existían, "
                f"{cont['ignoradas']} resúmenes ignorados, "
                f"{cont['saltadas_banco_subcat'] + cont['saltadas_monto']} salteadas, {cont['errores']} errores"
            )
    finally:
        wb.close()

    conn.execute("COMMIT;" if args.confirmar else "ROLLBACK;")
    db.desconectar()

    print(f"\n[FIN] NOE — {modo}")
    print(f"  Total hojas procesadas  : {hojas_ok}" + (f" (+{len(hojas_con_error)} con error, rollback: {', '.join(hojas_con_error)})" if hojas_con_error else ""))
    print(f"  Total filas procesadas  : {totales['total']}")
    print(f"  Resúmenes ignorados     : {totales['ignoradas']}")
    print(f"  Importadas              : {totales['importadas']}")
    print(f"  Ya existían             : {totales['ya_existian']}")
    print(f"  Saltadas (banco/subcat) : {totales['saltadas_banco_subcat']}")
    print(f"  Saltadas (monto 0/err)  : {totales['saltadas_monto']}")
    print(f"  Errores                 : {totales['errores']}")
    print(f"  Fechas corregidas       : {totales['fechas_corregidas']} (revisar las líneas ⚠️ con FECHA / AÑO)")
    if not args.confirmar:
        print("\nNada se escribió — revisá el reporte y corré de nuevo con --confirmar para aplicar.")


if __name__ == "__main__":
    main()
