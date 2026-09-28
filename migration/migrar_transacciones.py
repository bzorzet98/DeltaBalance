"""
DeltaBalance — migration/migrar_transacciones.py

Importa las transacciones de las hojas mensuales ("YYYY-MM", 2023-02 a
2026-09) de la planilla vieja (migration/sources/, ver su README.md) a
data/deltabalanceBZ.db. SQL directo vía DatabaseManager — no pasa por
TransactionService ni por ningún service (pedido explícito).

Pasos:
1. Categorías faltantes (CATEGORIAS_NUEVAS).
2. Cuentas destino de MAPA_BANCO que todavía no existan, + su
   cuentas_saldos en ARS. Normalmente ya las creó
   migrar_compras_cuotas.py; CAJA EFECTIVO no la crea ningún otro script.
3. Cada hoja mensual, en orden cronológico → transacciones.

--- Decisiones tomadas al relevar las hojas reales ---

- tipo_movimiento sale del SIGNO del monto en TODAS las filas ('ingreso'
  si es > 0, 'egreso' si es < 0), también en autotransferencias, cambio
  de moneda, inversiones y deuda — nunca 'movimiento'.
  vw_balance_cuentas (db/schema.sql) suma una fila 'movimiento' como
  POSITIVA (ELSE t.monto_minor) y monto_minor es siempre positivo: una
  transferencia saliente importada como 'movimiento' le SUMARÍA el monto
  a la cuenta de origen. Es el mismo criterio que la app:
  TransactionService.create_transfer() arma un egreso (origen) + un
  ingreso (destino) con la categoría Autotransferencia — lo que marca una
  fila como movimiento de capital es la categoría, no tipo_movimiento.
- Fecha: manda la hoja. Las hojas 2023-02 a 2023-12 tienen casi todas las
  fechas con el año SIGUIENTE (2024-MM: el mismo mes de la hoja, otro año
  — día/mes cargado sin año y Excel completó el año en curso), y hay
  casos sueltos parecidos en otras hojas ("15/12" cargado en la hoja de
  enero cae en diciembre del año de la hoja). Regla (_resolver_fecha()):
  si la fecha cae en el mes de la hoja, va tal cual; si no, se prueba la
  misma fecha con el año de la hoja (y con el anterior/siguiente, por el
  cruce diciembre/enero) y se toma la que caiga en el mes de la hoja o en
  uno contiguo; si ninguna cae cerca, o la fecha falta o no se entiende,
  va el día 1 del mes de la hoja (como pide la consigna). Toda corrección
  se avisa con la fecha original.
- Moneda: la columna MONEDA existe desde 2025; una hoja sin ella va toda
  en ARS. "CHL" (hojas 2025+) es el peso chileno mal tipeado: se trata
  como CLP (MONEDA_ALIAS), no como moneda desconocida → ARS, porque 50000
  CLP no son 50000 ARS. El monto se pasa a minor units con los decimales
  de la moneda (CLP 0, BTC 8), no siempre * 100.
- Idempotencia: la clave de la consigna (concepto + fecha + cuenta_id +
  monto_minor + moneda_id) más tipo_movimiento (un ajuste de +100 y otro
  de -100 el mismo día son dos filas distintas), CONTANDO APARICIONES: si
  la planilla tiene la misma transacción dos veces (dos cafés iguales el
  mismo día, misma cuenta), la segunda no es un duplicado. Se importan
  tantas como aparezcan, y una segunda corrida del script las encuentra
  todas y no agrega ninguna. Con la clave sola, la segunda se descartaría
  como "YA EXISTE" en la primera corrida.
- Cuentas y categorías se comparan por nombre ignorando mayúsculas y
  tildes (el seed tiene 'Cambio Moneda', 'Seguros', 'Autotransferencia'),
  y una que ya existe se reusa sin modificarla — mismo criterio que
  migrar_compras_cuotas.py. Una categoría se busca por subcategoría entre
  las activas, bajo cualquier categoría principal (SEGUROS ya existe como
  'EGRESOS FIJOS · Seguros': se reusa, no se duplica).
- Si una transacción viene en una moneda que su cuenta todavía no opera,
  se agrega esa moneda a cuentas_saldos (saldo inicial 0):
  vw_balance_cuentas une transacciones con cuentas_saldos por moneda, así
  que sin esa fila la transacción no entra en ningún saldo.
- Columnas: se detectan por nombre en la fila de encabezados de cada hoja
  (hay 4 formatos: con o sin columna N°, con o sin MONEDA, encabezado en
  la fila 1 o en la 2), tomando la PRIMERA aparición de cada nombre — las
  hojas de 2023 tienen a la derecha tablas de resumen (EGRESOS FIJOS,
  TARJETA ESTE MES, DEUDAS A COBRAR...) con sus propios CONCEPTO/PRECIO/
  FECHA, que no se migran. Los encabezados se comparan sin signos ni
  tildes (la hoja 2026-09 dice "CONCEPTO}"). El monto es la columna PRECIO.
- Trazabilidad (docs/MIGRATION_NOTES.md, principio 2): cada transacción
  lleva en `notas` "MIGRADO DESDE EXCEL — HOJA YYYY-MM, FILA N".
- Todo texto insertado va en MAYÚSCULAS. Los valores de enum que exige un
  CHECK del schema ('ingreso', 'egreso', 'debito', 'efectivo') quedan en
  minúscula, como los define db/schema.sql.

--- Transacción y seguridad ---

El setup (categorías y cuentas) y cada hoja van en su propia transacción:
si algo falla dentro de una hoja, se hace rollback de esa hoja sola y se
sigue con la siguiente. Una fila con datos que no sirven se saltea sin
escribir nada.

Mismo esquema que el resto de migration/: por default es DRY-RUN — todo
corre dentro de una transacción (cada hoja en un SAVEPOINT, para que el
rollback de una hoja no se lleve el setup), se imprime el reporte
completo y al final se hace rollback. Con --confirmar: backup de la DB,
commit del setup y commit de cada hoja al terminarla. Nunca corre contra
data/deltabalance.db (se rechaza aunque se pase por --db-path).
DatabaseManager.inicializar() sí crea la DB si no existe (schema + seed),
igual que al abrir la app.

Requiere openpyxl >= 3.1 (no está en environment.yml). Desde 3.1, una
celda con formato de fecha y un serial imposible (hay una en la hoja
2025-03) se lee como "#VALUE!" en vez de cortar la lectura de la hoja.
    pip install openpyxl

Uso:
    # 1) Dry-run (no escribe nada, solo reporta):
    python migration/migrar_transacciones.py

    # 2) Aplicar de verdad (hace backup antes):
    python migration/migrar_transacciones.py --confirmar

    # Opcional: contra otra DB (ej. una copia de prueba)
    python migration/migrar_transacciones.py --db-path /tmp/prueba.db --confirmar

Este script NO lo ejecuta Claude Code (CLAUDE.md §0.1) — lo corre el
usuario a mano.
"""

from __future__ import annotations

import argparse
import calendar
import re
import sqlite3
import sys
import unicodedata
import warnings
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    import openpyxl
    from openpyxl.utils.datetime import from_excel
except ImportError:
    sys.exit("❌ Falta openpyxl — instalalo en el entorno del proyecto:  pip install openpyxl")

from db.database import DatabaseManager, to_minor


# openpyxl avisa por cada extensión de Excel que no soporta (tablas
# dinámicas, validaciones) y por cada fecha imposible — no afecta la
# lectura de valores (las fechas imposibles se resuelven en _resolver_fecha()).
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

RAIZ = Path(__file__).resolve().parent.parent
DB_PATH = RAIZ / "data" / "deltabalanceBZ.db"
DB_PROTEGIDA = RAIZ / "data" / "deltabalance.db"
SOURCES_DIR = Path(__file__).resolve().parent / "sources"
# Mismo par de nombres que migrar_compras_cuotas.py — se usa el primero que exista.
XLSX_CANDIDATOS = [SOURCES_DIR / "REGISTRO PRINCIPAL.xlsx", SOURCES_DIR / "REGISTRO_PRINCIPAL.xlsx"]
PATRON_HOJA_MENSUAL = re.compile(r"(\d{4})-(\d{2})")
FILAS_BUSQUEDA_HEADER = 20
MONEDA_DEFAULT = "ARS"
# Código tipeado en la planilla → código real en la tabla monedas.
MONEDA_ALIAS = {"CHL": "CLP"}
NOTA_MIGRACION = "MIGRADO DESDE EXCEL — HOJA {hoja}, FILA {fila}"

# --- Columnas de cada hoja (fila de encabezados) ---
COL_CONCEPTO = "CONCEPTO"
COL_BANCO = "BANCO"
COL_SUBCATEGORIA = "SUBCATEGORIA"
COL_MONTO = "PRECIO"
COL_FECHA = "FECHA"
COL_MONEDA = "MONEDA"  # opcional: existe desde 2025
COLUMNAS_OBLIGATORIAS = (COL_CONCEPTO, COL_BANCO, COL_SUBCATEGORIA, COL_MONTO, COL_FECHA)

# --- Filas a ignorar ---
# Resúmenes del sistema anterior, no transacciones reales.
CONCEPTOS_IGNORAR = {
    "INGRESO FIJO", "EGRESO FIJO", "EGRESO VARIABLE", "INGRESO VARIABLE",
    "CREDITO ESTE MES", "CREDITO PROXIMO MES", "DEUDAS A COBRAR", "DEUDAS A PAGAR",
    "INGRESOS", "EGRESOS", "SALDO INICIAL", "DEUDAS", "INGRESO VAR",
    "INGRESO", "EGRESO",
}
# Filas de fórmulas del sistema anterior ("" = BANCO vacío).
BANCOS_IGNORAR = {"OTRO", "GENERAL", "ESTIMATIVO", ""}

# --- Mapeo de banco → cuenta ---
MAPA_BANCO = {
    "MP":           "MP",
    "NACION":       "NACION",
    "CREDICOOP":    "CREDICOOP",
    "GALICIA":      "GALICIA",
    "BILLETERA ER": "BILLETERA ER",
    "CAJA":         "CAJA EFECTIVO",
    "COCOS":        "COCOS",
    "BBVA":         "BBVA",
    "NARANJA X":    "NARANJA X",
    "NEXO":         "NEXO",
    "BUENBIT":      "BUENBIT",
    "BULL MARKET":  "BULL MARKET",
    "BRUBANK":      "BRUBANK",
    "FIWIND":       "FIWIND",
    "BINANCE":      "BINANCE",
    "MACRO":        "MACRO",
    "JOY":          "JOY",
    "SANTANDER":    "SANTANDER",
}
# Tipo con el que se crea una cuenta de MAPA_BANCO que todavía no existe
# (default 'debito', mismo tipo que les da migrar_compras_cuotas.py).
TIPO_CUENTA_NUEVA = {"CAJA EFECTIVO": "efectivo"}
TIPO_CUENTA_DEFAULT = "debito"

# --- Mapeo de subcategoría → categoría DB (None = ignorar la fila) ---
MAPA_SUBCATEGORIA: dict[str, Optional[str]] = {
    "SUELDO":                 "SUELDO",
    "INGRESO FIJO":           "SUELDO",
    "RENDIMIENTOS":           "RENDIMIENTOS",
    "COBRO DEUDA":            "COBRO DEUDA",
    "DEUDA NOE":              "COBRO DEUDA",
    "INGRESO VAR":            "INGRESO VARIABLE",
    "INGRESO VARIABLE":       "INGRESO VARIABLE",
    "INGRESO":                "INGRESO VARIABLE",
    "PROVEDURIA":             "ALIMENTOS",
    "COMIDAS Y BEBIDAS":      "ALIMENTOS",
    "ALCOHOL":                "GASTRONOMÍA",
    "SUPERMERCADO":           "SUPERMERCADO",
    "TRANSPORTE":             "TRANSPORTE / AUTO",
    "SEGUROS":                "SEGUROS",
    "OCIO":                   "OCIO",
    "PERSONAL":               "BIENESTAR Y DEPORTE",
    "REGALOS":                "REGALOS",
    "MASCOTAS":               "MASCOTAS",
    "HOGAR":                  "HOGAR",
    "SALUD":                  "SALUD",
    "VACACIONES":             "VACACIONES",
    "EDUCACION":              "EDUCACION",
    "EDUCACIÓN":              "EDUCACION",
    "VIVIENDA":               "ALQUILER / VIVIENDA",
    "SERVICIOS BÁSICOS":      "SERVICIOS",
    "CUENTAS Y SERVICIOS":    "SERVICIOS",
    "EGRESO FIJO":            "SERVICIOS",
    "COMPARTIDO":             "EGRESO VARIABLE",
    "SINC":                   "EGRESO VARIABLE",
    "EGRESO":                 "EGRESO VARIABLE",
    "EGRESO VAR":             "EGRESO VARIABLE",
    "EGRESO VARIABLE":        "EGRESO VARIABLE",
    "TRANSFERENCIA OTRO":     "EGRESO VARIABLE",
    "COMISIONES":             "EGRESO VARIABLE",
    "TARJETA":                "SERVICIOS",
    "TARJETA DE CREDITO":     "SERVICIOS",
    "DEVOLUCIÓN TARJETA":     "REINTEGRO",
    "REINTEGRO PROMOCION":    "REINTEGRO PROMOCION",
    "AUTOTRANSFERENCIA":      "AUTOTRANSFERENCIA",
    "TRANSFERENCIA":          "AUTOTRANSFERENCIA",
    "EXTRACCION":             "AUTOTRANSFERENCIA",
    "CAMBIO MONEDA":          "CAMBIO MONEDA",
    "CAMBIO MONEDA/RESERVAS": "CAMBIO MONEDA",
    "INVERSIONES":            "INVERSIONES",
    "DEUDA":                  "DEUDA",
    "RESERVA":                None,  # ignorar
}

# --- Categorías nuevas a crear si no existen ---
CATEGORIAS_NUEVAS = [
    ("EGRESOS VARIABLES",  "EGRESO VARIABLE",     "egreso"),
    ("EGRESOS VARIABLES",  "ALIMENTOS",           "egreso"),
    ("EGRESOS VARIABLES",  "GASTRONOMÍA",         "egreso"),
    ("EGRESOS VARIABLES",  "REGALOS",             "egreso"),
    ("EGRESOS VARIABLES",  "MASCOTAS",            "egreso"),
    ("EGRESOS VARIABLES",  "VACACIONES",          "egreso"),
    ("EGRESOS VARIABLES",  "SEGUROS",             "egreso"),
    ("EGRESOS VARIABLES",  "REINTEGRO PROMOCION", "egreso"),
    ("INGRESOS",           "INGRESO VARIABLE",    "ingreso"),
    ("INGRESOS",           "BECA",                "ingreso"),
    ("MOVIMIENTO CAPITAL", "CAMBIO MONEDA",       "movimiento"),
]


class MigracionError(Exception):
    """Problema estructural (falta la moneda default, una hoja sin sus columnas) — corta el script o la hoja."""


# ============================================================
# NORMALIZACIÓN
# ============================================================

def _texto(valor: Any) -> str:
    """strip + MAYÚSCULAS. None → ''."""
    if valor is None:
        return ""
    return str(valor).strip().upper()


def _clave(texto: str) -> str:
    """Clave de comparación: sin tildes, en mayúsculas, espacios colapsados ('Cambio Moneda' == 'CAMBIO MONEDA')."""
    sin_tildes = "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")
    return " ".join(sin_tildes.upper().split())


def _clave_header(texto: str) -> str:
    """Encabezado sin tildes ni signos sueltos: 'CONCEPTO}' → 'CONCEPTO', 'N°' → 'N'."""
    return " ".join(re.sub(r"[^A-Z0-9 /]", "", _clave(texto)).split())


def _parsear_fecha(valor: Any) -> Optional[date]:
    """datetime/date de openpyxl, serial de Excel o texto 'D/M/YYYY' / 'YYYY-MM-DD' → date."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):
        # Serial de Excel en una celda sin formato de fecha.
        try:
            return from_excel(valor).date()
        except (AttributeError, OverflowError, TypeError, ValueError):
            return None
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def _con_anio(fecha: date, anio: int) -> date:
    """Misma fecha en otro año (29/02 → 28/02 si ese año no es bisiesto)."""
    dia = min(fecha.day, calendar.monthrange(anio, fecha.month)[1])
    return date(anio, fecha.month, dia)


def _meses_de_distancia(fecha: date, anio: int, mes: int) -> int:
    return (fecha.year * 12 + fecha.month) - (anio * 12 + mes)


def _resolver_fecha(valor: Any, anio_hoja: int, mes_hoja: int) -> tuple[str, Optional[str]]:
    """
    (fecha 'YYYY-MM-DD', aviso o None). Ver docstring del módulo,
    "Fecha: manda la hoja".
    """
    primer_dia = date(anio_hoja, mes_hoja, 1)
    fecha = _parsear_fecha(valor)
    if fecha is None:
        if valor is None or str(valor).strip() == "":
            return primer_dia.isoformat(), f"SIN FECHA → {primer_dia.isoformat()}"
        return primer_dia.isoformat(), f"FECHA INVÁLIDA ({valor!r}) → {primer_dia.isoformat()}"

    if _meses_de_distancia(fecha, anio_hoja, mes_hoja) == 0:
        return fecha.isoformat(), None

    candidatas = [fecha] + [
        _con_anio(fecha, anio) for anio in (anio_hoja, anio_hoja - 1, anio_hoja + 1) if anio != fecha.year
    ]
    for distancia_aceptada in (0, 1):
        for candidata in candidatas:
            if abs(_meses_de_distancia(candidata, anio_hoja, mes_hoja)) == distancia_aceptada:
                if candidata == fecha:
                    return fecha.isoformat(), None  # mes contiguo, año bien: se respeta
                return candidata.isoformat(), f"AÑO CORREGIDO: {fecha.isoformat()} → {candidata.isoformat()}"

    return primer_dia.isoformat(), f"FECHA FUERA DEL MES DE LA HOJA ({fecha.isoformat()}) → {primer_dia.isoformat()}"


def _parsear_monto(valor: Any) -> Optional[float]:
    """Solo números de verdad: None, texto (incluidos '#N/A'/'#VALUE!') o bool → None."""
    if valor is None or isinstance(valor, bool) or not isinstance(valor, (int, float)):
        return None
    return float(valor)


# ============================================================
# SETUP: CATEGORÍAS Y CUENTAS
# ============================================================

def _categorias_activas(conn: sqlite3.Connection) -> list[sqlite3.Row]:
    return conn.execute(
        "SELECT id, categoria_principal, subcategoria, tipo FROM categorias WHERE activa = 1;"
    ).fetchall()


def _crear_categorias(conn: sqlite3.Connection) -> None:
    print("\n[CATEGORÍAS] Verificando categorías...")
    for principal, sub, tipo in CATEGORIAS_NUEVAS:
        existente = next(
            (c for c in _categorias_activas(conn) if _clave(c["subcategoria"]) == _clave(sub)), None,
        )
        if existente is not None:
            ubicacion = f"{existente['categoria_principal']} · {existente['subcategoria']}"
            como = "" if ubicacion == f"{principal} · {sub}" else f" (como '{ubicacion}' — se reusa, no se duplica)"
            print(f"  ⏭️  {sub} ya existe{como}")
            continue
        cur = conn.execute(
            "INSERT OR IGNORE INTO categorias (categoria_principal, subcategoria, tipo) VALUES (?, ?, ?);",
            (principal, sub, tipo),
        )
        if cur.rowcount == 0:
            # Existe con ese nombre exacto pero desactivada: no se revierte
            # un soft-delete hecho a propósito — las filas que la necesiten
            # quedan como error.
            print(f"  ⚠️  {principal} · {sub} existe DESACTIVADA — no se reactiva")
        else:
            print(f"  ✅ {sub} creada")


def _mapa_categorias(conn: sqlite3.Connection) -> dict[str, sqlite3.Row]:
    """Categorías activas por clave de subcategoría. Si una se repite bajo dos principales, gana la de tipo 'egreso'."""
    mapa: dict[str, sqlite3.Row] = {}
    for c in _categorias_activas(conn):
        clave = _clave(c["subcategoria"])
        previa = mapa.get(clave)
        if previa is None or (previa["tipo"] != "egreso" and c["tipo"] == "egreso"):
            mapa[clave] = c
    return mapa


def _asegurar_cuentas(conn: sqlite3.Connection, moneda_ars_id: int) -> dict[str, int]:
    """Nombre de cuenta (valor de MAPA_BANCO) → id. Crea las que falten; las que existen no se tocan."""
    print("\n[CUENTAS] Verificando cuentas...")
    ids: dict[str, int] = {}
    for nombre in sorted(set(MAPA_BANCO.values())):
        existentes = {_clave(r["nombre"]): r for r in conn.execute("SELECT id, nombre FROM cuentas;")}
        existente = existentes.get(_clave(nombre))
        if existente is not None:
            ids[nombre] = existente["id"]
            continue
        tipo = TIPO_CUENTA_NUEVA.get(nombre, TIPO_CUENTA_DEFAULT)
        conn.execute("INSERT OR IGNORE INTO cuentas (nombre, tipo) VALUES (?, ?);", (nombre, tipo))
        cuenta_id = conn.execute("SELECT id FROM cuentas WHERE nombre = ?;", (nombre,)).fetchone()["id"]
        conn.execute(
            "INSERT OR IGNORE INTO cuentas_saldos (cuenta_id, moneda_id, saldo_inicial_minor) VALUES (?, ?, 0);",
            (cuenta_id, moneda_ars_id),
        )
        ids[nombre] = cuenta_id
        print(f"  ✅ {nombre} creada ({tipo})")
    print(f"  {len(ids)} cuentas listas.")
    return ids


# ============================================================
# HOJAS MENSUALES
# ============================================================

def _encontrar_header(ws, hoja: str) -> tuple[int, dict[str, int]]:
    for numero, valores in enumerate(
        ws.iter_rows(min_row=1, max_row=FILAS_BUSQUEDA_HEADER, values_only=True), start=1,
    ):
        claves = [_clave_header(_texto(v)) for v in valores]
        if COL_CONCEPTO in claves and COL_BANCO in claves:
            columnas: dict[str, int] = {}
            for indice, clave in enumerate(claves):
                if clave:
                    columnas.setdefault(clave, indice)  # primera aparición = tabla principal
            faltantes = [c for c in COLUMNAS_OBLIGATORIAS if c not in columnas]
            if faltantes:
                raise MigracionError(f"A la hoja {hoja} le faltan columnas: {', '.join(faltantes)}.")
            return numero, columnas
    raise MigracionError(
        f"No se encontró la fila de encabezados (con {COL_CONCEPTO} y {COL_BANCO}) en las primeras "
        f"{FILAS_BUSQUEDA_HEADER} filas de la hoja {hoja}."
    )


def _contadores_vacios() -> dict[str, int]:
    return {
        "total": 0, "ignoradas": 0, "importadas": 0, "ya_existian": 0,
        "saltadas_banco_subcat": 0, "saltadas_monto": 0, "errores": 0, "fechas_corregidas": 0,
    }


def _procesar_hoja(
    conn: sqlite3.Connection,
    ws,
    hoja: str,
    ids_cuentas: dict[str, int],
    categorias: dict[str, sqlite3.Row],
    monedas: dict[str, sqlite3.Row],
    ocurrencias: dict[tuple, int],
) -> dict[str, int]:
    """
    Procesa una hoja completa (el caller maneja su transacción). `ocurrencias`
    acumula, para todo el run, cuántas veces apareció cada clave de
    idempotencia en la planilla — ver docstring del módulo.
    """
    cont = _contadores_vacios()
    anio_hoja, mes_hoja = (int(x) for x in PATRON_HOJA_MENSUAL.fullmatch(hoja).groups())
    mapa_banco = {_clave(k): v for k, v in MAPA_BANCO.items()}
    mapa_subcategoria = {_clave(k): v for k, v in MAPA_SUBCATEGORIA.items()}
    ignorar_concepto = {_clave(c) for c in CONCEPTOS_IGNORAR}
    ignorar_banco = {_clave(b) for b in BANCOS_IGNORAR}

    fila_header, columnas = _encontrar_header(ws, hoja)
    tiene_moneda = COL_MONEDA in columnas
    if not tiene_moneda and anio_hoja >= 2025:
        print(f"  ⚠️  La hoja {hoja} no tiene columna {COL_MONEDA} — todo en {MONEDA_DEFAULT}")

    for fila_excel, valores in enumerate(
        ws.iter_rows(min_row=fila_header + 1, values_only=True), start=fila_header + 1,
    ):
        def celda(col: str) -> Any:
            indice = columnas.get(col)
            return valores[indice] if indice is not None and indice < len(valores) else None

        concepto = _texto(celda(COL_CONCEPTO))
        if not concepto:
            continue  # fila vacía
        cont["total"] += 1
        etiqueta = f"[{fila_excel:03d}] {concepto}"

        if _clave(concepto) in ignorar_concepto:
            cont["ignoradas"] += 1
            continue

        # --- Banco ---
        banco = _texto(celda(COL_BANCO))
        if _clave(banco) in ignorar_banco:
            cont["saltadas_banco_subcat"] += 1
            continue
        nombre_cuenta = mapa_banco.get(_clave(banco))
        if nombre_cuenta is None:
            cont["saltadas_banco_subcat"] += 1
            print(f'  ⚠️  {etiqueta} | BANCO DESCONOCIDO: "{banco}" — salteando')
            continue
        cuenta_id = ids_cuentas[nombre_cuenta]

        # --- Subcategoría → categoría ---
        subcategoria = _texto(celda(COL_SUBCATEGORIA))
        if not subcategoria:
            cont["saltadas_banco_subcat"] += 1
            continue
        clave_sub = _clave(subcategoria)
        if clave_sub not in mapa_subcategoria:
            cont["saltadas_banco_subcat"] += 1
            print(f'  ⚠️  {etiqueta} | SUBCATEGORÍA DESCONOCIDA: "{subcategoria}" — salteando')
            continue
        destino = mapa_subcategoria[clave_sub]
        if destino is None:  # RESERVA
            cont["saltadas_banco_subcat"] += 1
            continue
        categoria = categorias.get(_clave(destino))
        if categoria is None:
            cont["errores"] += 1
            print(f"  ❌ {etiqueta} | LA CATEGORÍA DESTINO '{destino}' NO EXISTE (ACTIVA) EN LA DB — salteando")
            continue

        # --- Monto ---
        valor_monto = celda(COL_MONTO)
        monto = _parsear_monto(valor_monto)
        if monto is None or monto == 0:
            cont["saltadas_monto"] += 1
            if isinstance(valor_monto, str) and valor_monto.strip():
                print(f"  ⚠️  {etiqueta} | MONTO NO NUMÉRICO: {valor_monto!r} — salteando")
            continue

        avisos: list[str] = []

        # --- Fecha ---
        fecha, aviso_fecha = _resolver_fecha(celda(COL_FECHA), anio_hoja, mes_hoja)
        if aviso_fecha:
            cont["fechas_corregidas"] += 1
            avisos.append(aviso_fecha)

        # --- Moneda ---
        codigo_moneda = MONEDA_DEFAULT
        if tiene_moneda:
            moneda_txt = _texto(celda(COL_MONEDA))
            codigo_moneda = MONEDA_ALIAS.get(moneda_txt, moneda_txt)
            if _clave(codigo_moneda) not in monedas:
                avisos.append(f"MONEDA DESCONOCIDA: \"{moneda_txt or '(vacía)'}\" → {MONEDA_DEFAULT}")
                codigo_moneda = MONEDA_DEFAULT
        moneda = monedas[_clave(codigo_moneda)]

        tipo_movimiento = "ingreso" if monto > 0 else "egreso"
        monto_minor = abs(to_minor(monto, moneda["decimales"]))
        if monto_minor == 0:
            # Monto real pero menor a la unidad mínima de la moneda (ej. 0.3 CLP).
            cont["saltadas_monto"] += 1
            print(f"  ⚠️  {etiqueta} | MONTO {monto} REDONDEA A 0 EN {moneda['codigo']} — salteando")
            continue

        # --- Idempotencia (contando apariciones) ---
        clave_idem = (concepto, fecha, cuenta_id, monto_minor, moneda["id"], tipo_movimiento)
        ocurrencias[clave_idem] = ocurrencias.get(clave_idem, 0) + 1
        en_db = conn.execute(
            """
            SELECT COUNT(*) AS n FROM transacciones
            WHERE concepto = ? AND fecha = ? AND cuenta_id = ? AND monto_minor = ?
              AND moneda_id = ? AND tipo_movimiento = ?;
            """,
            clave_idem,
        ).fetchone()["n"]
        if en_db >= ocurrencias[clave_idem]:
            cont["ya_existian"] += 1
            print(f"  ⏭️  {etiqueta} | YA EXISTE")
            continue

        # --- La cuenta tiene que operar en la moneda de la transacción ---
        cur = conn.execute(
            "INSERT OR IGNORE INTO cuentas_saldos (cuenta_id, moneda_id, saldo_inicial_minor) VALUES (?, ?, 0);",
            (cuenta_id, moneda["id"]),
        )
        if cur.rowcount:
            avisos.append(f"{nombre_cuenta} AHORA OPERA TAMBIÉN EN {moneda['codigo']} (cuentas_saldos)")

        conn.execute(
            """
            INSERT INTO transacciones
                (fecha, concepto, cuenta_id, categoria_id, moneda_id, tipo_movimiento, monto_minor, notas)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?);
            """,
            (
                fecha, concepto, cuenta_id, categoria["id"], moneda["id"], tipo_movimiento, monto_minor,
                NOTA_MIGRACION.format(hoja=hoja, fila=fila_excel),
            ),
        )
        cont["importadas"] += 1
        for aviso in avisos:
            print(f"  ⚠️  {etiqueta} | {aviso}")
        print(
            f"  ✅ {etiqueta} | {nombre_cuenta} | {moneda['codigo']} ${abs(monto):,.2f} | "
            f"{tipo_movimiento} | {fecha}"
        )
    return cont


# ============================================================
# MAIN
# ============================================================

def _mostrar_ruta(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(RAIZ))
    except ValueError:
        return str(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Migra las transacciones de las hojas mensuales de la planilla Excel.")
    parser.add_argument("--db-path", default=None, help=f"Ruta a la DB (default: {_mostrar_ruta(DB_PATH)}).")
    parser.add_argument("--confirmar", action="store_true", help="Aplica los cambios. Sin esto, solo dry-run.")
    
    parser.add_argument('--desde', default=None, 
        help='Procesar solo hojas desde este mes (ej: 2026-01)')

    args = parser.parse_args()

    db_path = Path(args.db_path) if args.db_path else DB_PATH
    if db_path.resolve() == DB_PROTEGIDA.resolve():
        sys.exit(f"❌ Este script nunca corre contra {_mostrar_ruta(DB_PROTEGIDA)} (la DB real de la app).")

    xlsx_path = next((p for p in XLSX_CANDIDATOS if p.exists()), None)
    if xlsx_path is None:
        buscados = ", ".join(f"'{p.name}'" for p in XLSX_CANDIDATOS)
        sys.exit(f"❌ No se encontró la planilla en {_mostrar_ruta(SOURCES_DIR)} (se buscó {buscados}). Ver su README.md.")

    modo = "APLICANDO CAMBIOS" if args.confirmar else "DRY-RUN (nada se escribe)"
    print(f"[INICIO] Migración de transacciones mensuales — {modo}")
    print(f"[XLSX] {_mostrar_ruta(xlsx_path)}")
    print(f"[DB] Conectando a {_mostrar_ruta(db_path)}")

    db = DatabaseManager(db_path)
    db.inicializar()  # idempotente: schema + migraciones de columna; seed solo si la DB es nueva
    conn = db.conn
    # inicializar() puede dejar abierta una transacción implícita (el
    # backfill de una migración de columna recién aplicada no commitea) —
    # se cierra acá, antes de pasar a control manual de transacciones.
    conn.commit()
    # Control manual (BEGIN/SAVEPOINT/COMMIT explícitos): cada hoja es un
    # SAVEPOINT dentro de la transacción abierta, así su rollback no se
    # lleva el setup ni las hojas anteriores en el dry-run.
    conn.isolation_level = None

    if args.confirmar:
        # El backup copia solo el archivo .db — con journal_mode=WAL, lo
        # que siga en el -wal no estaría en la copia sin este checkpoint.
        conn.execute("PRAGMA wal_checkpoint(TRUNCATE);")
        db.hacer_backup()

    conn.execute("BEGIN;")
    try:
        monedas = {_clave(m["codigo"]): m for m in conn.execute("SELECT id, codigo, decimales FROM monedas;")}
        if _clave(MONEDA_DEFAULT) not in monedas:
            raise MigracionError(f"La moneda default {MONEDA_DEFAULT} no existe en la DB.")
        _crear_categorias(conn)
        categorias = _mapa_categorias(conn)
        ids_cuentas = _asegurar_cuentas(conn, monedas[_clave(MONEDA_DEFAULT)]["id"])

        destinos = {d for d in MAPA_SUBCATEGORIA.values() if d is not None}
        faltantes = sorted(d for d in destinos if _clave(d) not in categorias)
        if faltantes:
            print(f"\n  ⚠️  Categorías destino que NO existen activas en la DB (sus filas van a Errores): {', '.join(faltantes)}")
    except Exception as err:
        conn.execute("ROLLBACK;")
        db.desconectar()
        print(f"\n❌ ERROR en el setup — rollback completo, no se escribió nada: {err}")
        if isinstance(err, MigracionError):
            sys.exit(1)
        raise

    if args.confirmar:
        conn.execute("COMMIT;")
        conn.execute("BEGIN;")

    totales = _contadores_vacios()
    hojas_ok = 0
    hojas_con_error: list[str] = []
    ocurrencias: dict[tuple, int] = {}

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        hojas = sorted(n for n in wb.sheetnames if PATRON_HOJA_MENSUAL.fullmatch(n))
        print(f"\n{len(hojas)} hojas mensuales: {hojas[0] if hojas else '-'} … {hojas[-1] if hojas else '-'}")
        for hoja in hojas:
            if args.desde and hoja < args.desde:
                continue
            print(f"\n[HOJA {hoja}]")
            ocurrencias_antes = dict(ocurrencias)
            conn.execute("SAVEPOINT hoja;")
            try:
                cont = _procesar_hoja(conn, wb[hoja], hoja, ids_cuentas, categorias, monedas, ocurrencias)
            except Exception as err:
                conn.execute("ROLLBACK TO hoja;")
                conn.execute("RELEASE hoja;")
                ocurrencias = ocurrencias_antes
                hojas_con_error.append(hoja)
                print(f"  ❌ ERROR en la hoja {hoja} — rollback de la hoja, se sigue con la siguiente: {err!r}")
                continue
            conn.execute("RELEASE hoja;")
            if args.confirmar:
                conn.execute("COMMIT;")
                conn.execute("BEGIN;")
            hojas_ok += 1
            for k, v in cont.items():
                totales[k] += v
            print(
                f"  → {cont['total']} filas: {cont['importadas']} importadas, {cont['ya_existian']} ya existían, "
                f"{cont['ignoradas']} resúmenes ignorados, "
                f"{cont['saltadas_banco_subcat'] + cont['saltadas_monto']} salteadas, {cont['errores']} errores"
            )
    finally:
        wb.close()

    conn.execute("COMMIT;" if args.confirmar else "ROLLBACK;")
    db.desconectar()

    print(f"\n[FIN] {modo}")
    print(f"  Total hojas procesadas  : {hojas_ok}" + (f" (+{len(hojas_con_error)} con error, rollback: {', '.join(hojas_con_error)})" if hojas_con_error else ""))
    print(f"  Total filas procesadas  : {totales['total']}")
    print(f"  Resúmenes ignorados     : {totales['ignoradas']}")
    print(f"  Importadas              : {totales['importadas']}")
    print(f"  Ya existían             : {totales['ya_existian']}")
    print(f"  Saltadas (banco/subcat) : {totales['saltadas_banco_subcat']}")
    print(f"  Saltadas (monto 0/err)  : {totales['saltadas_monto']}")
    print(f"  Errores                 : {totales['errores']}")
    print(f"  Fechas corregidas       : {totales['fechas_corregidas']} (revisar las líneas ⚠️ con FECHA / AÑO)")
    if not args.confirmar:
        print("\nNada se escribió — revisá el reporte y corré de nuevo con --confirmar para aplicar.")


if __name__ == "__main__":
    main()
