"""
DeltaBalance — migration/migrar_deudas_noe.py

Importa las deudas de la planilla de Noe (migration/sources/Finanzas.xlsx)
al libro de deudas de data/deltabalanceNR.db (estructura final de `deudas`,
docs/DATA_MODEL_DECISIONS.md sección 22: dos tabs, monto_minor CON SIGNO —
positivo = la deuda crece, negativo = un pago). Basado en
migration/migrar_deudas.py (el de Bruno), que lee un CSV; este lee el Excel
directo, con dos hojas:

    TABLA DEUDAS A COBRAR → tab 'me_deben'
        CONCEPTO, PERSONA, TIPO DE DEUDA, PRECIO, [columna negativa], FECHA, MONEDA
    TABLA DEUDAS A PAGAR  → tab 'debo'
        CONCEPTO, PERSONA, CATEGORIA, PRECIO, [columna negativa], FECHA, MONEDA

Encabezado en la fila 2 (se busca en las primeras FILAS_BUSQUEDA_HEADER
filas, por si se corre). Las columnas se encuentran por NOMBRE; la columna
negativa es la que sigue a PRECIO, se llame como se llame.

Reglas (pedido explícito, las mismas en las dos hojas):
- PRECIO positivo → +PRECIO (la deuda crece).
- Columna negativa con valor → −|valor| (un pago: recibido en ME DEBEN,
  hecho en DEBO).
- Las dos columnas con valor → dos filas.
- PRECIO negativo (no lo cubre el pedido) → −|PRECIO|, como un pago — mismo
  criterio que el script de Bruno; se avisa cuántas.
- Se saltean las filas sin PERSONA o sin ningún monto (las filas
  completamente vacías ni se cuentan).
- Moneda vacía → ARS (se informa cuántas); un código que no está en la
  tabla monedas → la fila se saltea y se informa.
- Persona: sin espacios de más y en mayúsculas (utils/personas.py, la misma
  regla que DebtsService).
- Monto: un número de la celda va tal cual; un texto pasa por _numero() de
  migrar_deudas.py (formato argentino: "1.234,56", "$ 500", "(500)"). Minor
  units con los decimales de la moneda (ARS/USD: × 100).
- Fecha: datetime de openpyxl, serial de Excel o texto DD/MM/AAAA (y los
  demás formatos de migrar_deudas._fecha()) → AAAA-MM-DD.
- TIPO DE DEUDA / CATEGORIA no cambian nada (a diferencia del COBRO DEUDA
  del script de Bruno): el reporte lista sus valores para revisarlos.
- Notas: "MIGRADO DESDE EXCEL NOE — <HOJA>, FILA N" (N = fila del Excel).
  Lleva el nombre de la hoja y no solo "TABLA DEUDAS": la FILA 5 de A
  COBRAR y la FILA 5 de A PAGAR tienen que ser notas distintas, porque la
  nota es la clave de idempotencia — una fila cuyas notas ya están en
  `deudas` se saltea (se puede correr de nuevo sin duplicar).

Se inserta con SQL directo, no con DebtsService: su chequeo de doble-click
rechazaría dos movimientos idénticos legítimos del Excel cargados uno tras
otro (mismo criterio que el script de Bruno).

Por default es DRY-RUN: importa sobre una COPIA temporal de la base, imprime
el reporte y descarta la copia — la base real no se toca. Con --confirmar:
inicializar(), checkpoint + backup (DatabaseManager.hacer_backup()) e
importación de las dos hojas en una sola transacción. Nunca corre contra
data/deltabalance.db ni contra la base de Bruno (data/deltabalanceBZ.db),
aunque se pase por --db-path.

Uso:
    # 1) Dry-run (no escribe nada en la base real, solo reporta):
    python migration/migrar_deudas_noe.py

    # 2) Aplicar de verdad (hace backup antes):
    python migration/migrar_deudas_noe.py --confirmar

    # Opcional: contra otra DB (ej. una copia de prueba)
    python migration/migrar_deudas_noe.py --db-path /tmp/prueba_noe.db --confirmar

Después de aplicar, conviene recalcular los snapshots (botón ↻ de Deudas)
para que el SALDO ANTERIOR incluya lo importado.

Este script NO lo ejecuta Claude Code (CLAUDE.md §0.1) — lo corre el
usuario a mano.
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
import tempfile
import warnings
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    import openpyxl
    from openpyxl.utils.datetime import from_excel
except ImportError:
    sys.exit("❌ Falta openpyxl — instalalo en el entorno del proyecto:  pip install openpyxl")

from db.database import DatabaseManager
from migration.migrar_deudas import _clave, _numero
from migration.migrar_deudas import _fecha as _fecha_texto
from utils.personas import normalizar_persona

# openpyxl avisa por cada extensión de Excel que no soporta — no afecta la lectura de valores.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

RAIZ = Path(__file__).resolve().parent.parent
DB_PATH = RAIZ / "data" / "deltabalanceNR.db"
# La DB real de la app y la de Bruno: este script nunca las toca.
DBS_PROHIBIDAS = (RAIZ / "data" / "deltabalance.db", RAIZ / "data" / "deltabalanceBZ.db")
XLSX_PATH = RAIZ / "migration" / "sources" / "Finanzas.xlsx"
NOTA = "MIGRADO DESDE EXCEL NOE — {hoja}, FILA {fila}"
MONEDA_SI_VACIA = "ARS"
FILAS_BUSQUEDA_HEADER = 10

# Columnas por nombre (la negativa: la que sigue a PRECIO).
COL_CONCEPTO = "CONCEPTO"
COL_PERSONA = "PERSONA"
COL_PRECIO = "PRECIO"
COL_FECHA = "FECHA"
COL_MONEDA = "MONEDA"
COLUMNAS_OBLIGATORIAS = (COL_CONCEPTO, COL_PERSONA, COL_PRECIO, COL_FECHA, COL_MONEDA)


@dataclass(frozen=True)
class Hoja:
    nombre: str
    tab: str
    etiqueta_tab: str
    columna_tipo: str  # solo para el reporte (ver docstring)


HOJAS = (
    Hoja("TABLA DEUDAS A COBRAR", "me_deben", "ME DEBEN", "TIPO DE DEUDA"),
    Hoja("TABLA DEUDAS A PAGAR", "debo", "DEBO", "CATEGORIA"),
)


class MigracionError(Exception):
    """Problema estructural (falta la hoja o una columna) — corta el script antes de tocar la base."""


@dataclass
class ReporteHoja:
    filas_datos: int = 0
    insertadas: Counter = field(default_factory=Counter)  # "positivas" / "negativas"
    salteadas: Counter = field(default_factory=Counter)
    avisos: Counter = field(default_factory=Counter)
    tipos: Counter = field(default_factory=Counter)
    errores: list[str] = field(default_factory=list)


# ============================================================
# LECTURA DEL EXCEL
# ============================================================

def _celda_texto(valor: Any) -> str:
    return "" if valor is None else str(valor).strip()


def _monto(valor: Any) -> Optional[float]:
    """Un número de openpyxl va tal cual; un texto, por _numero() (formato argentino). None si no hay monto."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    return _numero(str(valor))


def _fecha(valor: Any) -> Optional[str]:
    """datetime/date de openpyxl, serial de Excel o texto (migrar_deudas._fecha()) → 'YYYY-MM-DD'."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, (int, float)):
        try:
            return from_excel(valor).date().isoformat()
        except (AttributeError, OverflowError, TypeError, ValueError):
            return None
    return _fecha_texto(str(valor))


def _leer_hoja(wb, hoja: Hoja) -> list[dict[str, Any]]:
    """Las filas de datos de la hoja: {fila, concepto, persona, tipo, precio, negativa, fecha, moneda} (valores crudos)."""
    if hoja.nombre not in wb.sheetnames:
        raise MigracionError(f"La planilla no tiene la hoja '{hoja.nombre}'.")
    ws = wb[hoja.nombre]
    filas_header = ws.iter_rows(min_row=1, max_row=FILAS_BUSQUEDA_HEADER, values_only=True)
    for numero_header, valores in enumerate(filas_header, start=1):
        claves = [_clave(_celda_texto(v)) for v in valores]
        if COL_CONCEPTO in claves and COL_PERSONA in claves:
            break
    else:
        raise MigracionError(
            f"No se encontró la fila de encabezados (con {COL_CONCEPTO} y {COL_PERSONA}) en las primeras "
            f"{FILAS_BUSQUEDA_HEADER} filas de '{hoja.nombre}'."
        )
    columnas: dict[str, int] = {}
    for indice, clave in enumerate(claves):
        if clave:
            columnas.setdefault(clave, indice)  # primera aparición
    faltantes = [c for c in COLUMNAS_OBLIGATORIAS if c not in columnas]
    if faltantes:
        raise MigracionError(f"A la hoja '{hoja.nombre}' le faltan columnas: {', '.join(faltantes)}.")
    indice_negativa = columnas[COL_PRECIO] + 1
    if indice_negativa < len(claves) and claves[indice_negativa] in COLUMNAS_OBLIGATORIAS:
        raise MigracionError(
            f"En '{hoja.nombre}' la columna después de {COL_PRECIO} es {claves[indice_negativa]}, "
            "no la columna negativa (ver docstring)."
        )
    indices = {
        "concepto": columnas[COL_CONCEPTO],
        "persona": columnas[COL_PERSONA],
        "tipo": columnas.get(_clave(hoja.columna_tipo)),
        "precio": columnas[COL_PRECIO],
        "negativa": indice_negativa,
        "fecha": columnas[COL_FECHA],
        "moneda": columnas[COL_MONEDA],
    }

    filas: list[dict[str, Any]] = []
    for numero_fila, valores in enumerate(
        ws.iter_rows(min_row=numero_header + 1, values_only=True), start=numero_header + 1,
    ):
        fila = {
            nombre: (valores[i] if i is not None and i < len(valores) else None) for nombre, i in indices.items()
        }
        if all(_celda_texto(v) == "" for v in fila.values()):
            continue  # fila vacía (formato aplicado más abajo de los datos)
        fila["fila"] = numero_fila
        filas.append(fila)
    return filas


# ============================================================
# IMPORTACIÓN
# ============================================================

def _movimientos(fila: dict[str, Any], reporte: ReporteHoja) -> list[float]:
    """Los montos CON SIGNO de una fila (ver reglas en el docstring): 0, 1 o 2."""
    movimientos: list[float] = []
    precio = _monto(fila["precio"])
    if precio:
        if precio < 0:
            reporte.avisos["PRECIO negativo → pago (negativo)"] += 1
        movimientos.append(-abs(precio) if precio < 0 else abs(precio))
    negativa = _monto(fila["negativa"])
    if negativa:
        movimientos.append(-abs(negativa))
    if len(movimientos) == 2:
        reporte.avisos["con las dos columnas → dos filas"] += 1
    return movimientos


def _importar(conn: sqlite3.Connection, hoja: Hoja, filas: list[dict[str, Any]]) -> ReporteHoja:
    """Inserta las filas de una hoja en `deudas` (sin commit: decide el caller)."""
    reporte = ReporteHoja(filas_datos=len(filas))
    monedas = {_clave(f["codigo"]): f for f in conn.execute("SELECT id, codigo, decimales FROM monedas;")}
    ya_importadas = {
        f["notas"] for f in conn.execute("SELECT notas FROM deudas WHERE notas LIKE ?;", (f"{NOTA.format(hoja=hoja.nombre, fila='')}%",))
    }

    for fila in filas:
        tipo = _celda_texto(fila["tipo"]).upper()
        reporte.tipos[tipo or "(vacío)"] += 1
        persona = normalizar_persona(_celda_texto(fila["persona"]))
        if not persona:
            reporte.salteadas["sin PERSONA"] += 1
            continue
        nota = NOTA.format(hoja=hoja.nombre, fila=fila["fila"])
        if nota in ya_importadas:
            reporte.salteadas["ya importada (misma nota)"] += 1
            continue
        movimientos = _movimientos(fila, reporte)
        if not movimientos:
            reporte.salteadas["sin monto"] += 1
            continue

        codigo = _clave(_celda_texto(fila["moneda"]))
        if not codigo:
            reporte.avisos[f"MONEDA vacía → {MONEDA_SI_VACIA}"] += 1
            codigo = MONEDA_SI_VACIA
        moneda = monedas.get(codigo)
        if moneda is None:
            reporte.salteadas[f"moneda desconocida ({codigo})"] += 1
            continue
        fecha = _fecha(fila["fecha"])
        if fecha is None:
            reporte.errores.append(f"fila {fila['fila']}: fecha inválida {fila['fecha']!r} — salteada")
            continue

        concepto = _celda_texto(fila["concepto"]) or None
        for valor in movimientos:
            monto_minor = round(valor * 10 ** moneda["decimales"])
            if monto_minor == 0:
                reporte.errores.append(f"fila {fila['fila']}: monto {valor!r} redondea a 0 — salteado")
                continue
            conn.execute(
                """
                INSERT INTO deudas (entidad_persona, concepto, tab, monto_minor, moneda_id, fecha, notas, origen_tipo)
                VALUES (?, ?, ?, ?, ?, ?, ?, 'manual');
                """,
                (persona, concepto, hoja.tab, monto_minor, moneda["id"], fecha, nota),
            )
            reporte.insertadas["positivas" if monto_minor > 0 else "negativas"] += 1
    return reporte


def _netos(conn: sqlite3.Connection, tab: str) -> list[sqlite3.Row]:
    """Saldo neto por (persona, moneda) del tab en la base, ya con lo importado; sin los saldados (0)."""
    return conn.execute(
        """
        SELECT d.entidad_persona, m.codigo, m.simbolo, m.decimales, SUM(d.monto_minor) AS neto
        FROM deudas d JOIN monedas m ON m.id = d.moneda_id
        WHERE d.tab = ?
        GROUP BY d.entidad_persona, d.moneda_id
        HAVING neto != 0
        ORDER BY d.entidad_persona, m.codigo;
        """,
        (tab,),
    ).fetchall()


def _importar_todo(conn: sqlite3.Connection, filas_por_hoja: dict[str, list[dict[str, Any]]]) -> tuple[dict, dict]:
    """Las dos hojas en la transacción abierta del caller → (reporte por hoja, netos por tab)."""
    reportes = {hoja.nombre: _importar(conn, hoja, filas_por_hoja[hoja.nombre]) for hoja in HOJAS}
    netos = {hoja.tab: _netos(conn, hoja.tab) for hoja in HOJAS}
    return reportes, netos


# ============================================================
# REPORTE
# ============================================================

def _mostrar_ruta(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(RAIZ))
    except ValueError:
        return str(path)


def _imprimir_reporte(reportes: dict[str, ReporteHoja], netos: dict[str, list[sqlite3.Row]], modo: str) -> None:
    for hoja in HOJAS:
        r = reportes[hoja.nombre]
        print(f"\n[{hoja.nombre} → {hoja.etiqueta_tab}] {r.filas_datos} fila(s) de datos")
        print(f"  ✅ importadas: {sum(r.insertadas.values())} ({r.insertadas['positivas']} positivas, "
              f"{r.insertadas['negativas']} negativas = pagos)")
        for motivo, cantidad in sorted(r.salteadas.items()):
            print(f"  ⏭️  salteadas — {motivo}: {cantidad}")
        for aviso, cantidad in sorted(r.avisos.items()):
            print(f"  ⚠️  {aviso}: {cantidad}")
        for error in r.errores:
            print(f"  ❌ {error}")
        tipos = ", ".join(f"{tipo} ({cantidad})" for tipo, cantidad in sorted(r.tipos.items()))
        print(f"  {hoja.columna_tipo} (no cambia nada, para revisar): {tipos or '—'}")

    ancho_hoja = max(len(h.nombre) for h in HOJAS) + 1
    print(f"\n[FIN] NOE — {modo}")
    for hoja in HOJAS:
        print(f"  {hoja.nombre + ':':<{ancho_hoja}} {sum(reportes[hoja.nombre].insertadas.values())} filas → {hoja.tab}")
    print(f"  Total importadas: {sum(sum(r.insertadas.values()) for r in reportes.values())}")
    print(f"  Saltadas: {sum(sum(r.salteadas.values()) for r in reportes.values())}")
    print(f"  Errores: {sum(len(r.errores) for r in reportes.values())}")

    for hoja in HOJAS:
        print(f"\n[NETO POR PERSONA — {hoja.etiqueta_tab}] (en la base, ya con lo importado)")
        if not netos[hoja.tab]:
            print("  —")
        ancho = max((len(n["entidad_persona"]) for n in netos[hoja.tab]), default=0)
        for n in netos[hoja.tab]:
            signo = "-" if n["neto"] < 0 else "+"
            monto = abs(n["neto"]) / 10 ** n["decimales"]
            print(f"  {n['entidad_persona']:<{ancho}}   {signo}{n['simbolo'] or ''}{monto:,.{n['decimales']}f}  {n['codigo']}")


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    parser = argparse.ArgumentParser(description="Importa las deudas de la planilla de Noe al libro de deudas.")
    parser.add_argument("--db-path", default=None, help=f"Ruta a la DB (default: {_mostrar_ruta(DB_PATH)}).")
    parser.add_argument("--confirmar", action="store_true", help="Aplica los cambios. Sin esto, solo dry-run.")
    args = parser.parse_args()

    db_path = Path(args.db_path) if args.db_path else DB_PATH
    for prohibida in DBS_PROHIBIDAS:
        if db_path.resolve() == prohibida.resolve():
            sys.exit(f"❌ Este script nunca corre contra {_mostrar_ruta(prohibida)} (es de Noe: {_mostrar_ruta(DB_PATH)}).")
    if not XLSX_PATH.exists():
        sys.exit(f"❌ No se encontró la planilla de Noe: {_mostrar_ruta(XLSX_PATH)}.")
    if not db_path.exists():
        sys.exit(f"❌ No existe la base {_mostrar_ruta(db_path)}.")

    modo = "APLICANDO CAMBIOS" if args.confirmar else "DRY-RUN"
    print(f"[INICIO] Migración de deudas de NOE — {modo if args.confirmar else 'DRY-RUN (la base real no se toca)'}")
    print(f"[XLSX] {_mostrar_ruta(XLSX_PATH)}")
    print(f"[DB] {_mostrar_ruta(db_path)}")

    # El Excel se lee entero antes de tocar la base: una hoja o columna que falta corta acá.
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    try:
        filas_por_hoja = {hoja.nombre: _leer_hoja(wb, hoja) for hoja in HOJAS}
    except MigracionError as err:
        sys.exit(f"❌ {err}")
    finally:
        wb.close()

    if not args.confirmar:
        with tempfile.TemporaryDirectory(prefix="deltabalance_deudas_noe_") as carpeta:
            copia = Path(carpeta) / "copia.db"
            origen, destino = sqlite3.connect(db_path), sqlite3.connect(copia)
            try:
                origen.backup(destino)  # copia consistente, incluido lo que esté en el -wal
            finally:
                destino.close()
                origen.close()
            db = DatabaseManager(copia)
            try:
                db.inicializar()
                conn = db.conn
                conn.commit()
                conn.execute("BEGIN;")
                reportes, netos = _importar_todo(conn, filas_por_hoja)
                conn.rollback()
            finally:
                db.desconectar()
        _imprimir_reporte(reportes, netos, modo)
        print("\nNada se escribió en la base real — revisá el reporte y corré de nuevo con --confirmar para aplicar.")
        return

    db = DatabaseManager(db_path)
    try:
        db.inicializar()
        conn = db.conn
        # inicializar() puede dejar abierta una transacción implícita (backfill de una migración de columna).
        conn.commit()
        # El backup copia solo el archivo .db — con WAL, lo que siga en el -wal no estaría sin este checkpoint.
        conn.execute("PRAGMA wal_checkpoint(TRUNCATE);")
        db.hacer_backup()
        conn.execute("BEGIN;")
        try:
            reportes, netos = _importar_todo(conn, filas_por_hoja)
            conn.commit()
        except Exception:
            conn.rollback()
            print("\n❌ ERROR — rollback completo, no se importó nada.")
            raise
    finally:
        db.desconectar()
    _imprimir_reporte(reportes, netos, modo)
    print("\n✅ Deudas importadas. Recalculá los snapshots (↻ en Deudas) para que el SALDO ANTERIOR las incluya.")


if __name__ == "__main__":
    main()
