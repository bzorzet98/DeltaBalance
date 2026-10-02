"""
DeltaBalance — migration/migrar_compras_cuotas_noe.py

Importa TODAS las compras en cuotas de la hoja "REGISTRO COMPRA EN CUOTAS"
de la planilla de Noe (migration/sources/Finanzas.xlsx) a
data/deltabalanceNR.db. Misma mecánica que migration/migrar_compras_cuotas.py
(el de Bruno): SQL directo vía DatabaseManager, sin pasar por FeesService
ni por ningún service, con el cronograma replicado acá. Lo que cambia:

- Cuentas: BERSA (débito) y dos tarjetas de crédito — BERSA VISA EXT RP
  (extensión de la tarjeta madre, se paga desde BERSA) y BZ (compras de
  Bruno con su tarjeta: sin cuenta débito asociada, cuenta_pago_id NULL).
- Banco: VISA → BERSA VISA EXT RP y OTRO → BZ. OTRO NO se saltea (en el de
  Bruno sí): acá son las compras hechas con la tarjeta de Bruno.
- Categoría: todo va a EGRESO VARIABLE — TARJETA por MAPA_CATEGORIA y
  cualquier otra subcategoría por fallback, sin aviso por fila (pedido
  explícito: es el comportamiento esperado). No hay inferencia por
  concepto. El reporte final cuenta cuántas fueron por fallback.
  EGRESO VARIABLE se crea si no existe (misma categoría principal y tipo
  que en el script de Bruno).
- Notas: "MIGRADO DESDE EXCEL NOE — COMPRAS EN CUOTAS, FILA N".
- Columna de la cantidad de cuotas: la planilla de Noe dice "CANTIDAD DE
  CUOTAS" y la de Bruno "CANTIDAD DE CUOTAS TOTALES" — se aceptan las dos
  (ALIAS_CUOTAS_TOTALES, _encontrar_header() propio).

Igual que en el de Bruno — sus funciones se importan, no se copian:
detección de columnas, parseo de fechas / montos / cantidades, cronograma
desde FECHA PRIMER PAGO (PROYECTAR_DESDE_PRIMER_PAGO), las primeras
CANTIDAD DE CUOTAS PAGADAS cuotas en 'pagado' (MARCAR_CUOTAS_YA_PAGADAS),
montos <= 0 tal cual (reintegros), moneda nueva → cuentas_saldos de la
tarjeta, cuentas reusadas por nombre sin modificarlas, e idempotencia por
concepto + fecha_compra + cuenta_id + monto_total_minor. El porqué de cada
regla está en el docstring de migrar_compras_cuotas.py.

--- Transacción y seguridad ---

Todo (cuentas, categoría y filas) va en UNA sola transacción. Una fila con
datos inválidos se saltea y cuenta como error sin escribir nada; un error
de base de datos (o de otro tipo) hace rollback de TODO y corta el script.
Por default es DRY-RUN: corre todo, imprime el reporte y hace rollback. Con
--confirmar hace backup de la DB y commitea. Nunca corre contra
data/deltabalance.db ni contra la base de Bruno (data/deltabalanceBZ.db),
aunque se pase por --db-path. DatabaseManager.inicializar() crea la DB si
no existe (schema + seed), igual que al abrir la app.

Requiere openpyxl (pip install openpyxl).

Uso:
    # 1) Dry-run (no escribe nada, solo reporta):
    python migration/migrar_compras_cuotas_noe.py

    # 2) Aplicar de verdad (hace backup antes):
    python migration/migrar_compras_cuotas_noe.py --confirmar

    # Opcional: contra otra DB (ej. una copia de prueba)
    python migration/migrar_compras_cuotas_noe.py --db-path /tmp/prueba_noe.db --confirmar

Este script NO lo ejecuta Claude Code (CLAUDE.md §0.1) — lo corre el
usuario a mano.
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
from pathlib import Path
from typing import Any, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    import openpyxl
except ImportError:
    sys.exit("❌ Falta openpyxl — instalalo en el entorno del proyecto:  pip install openpyxl")

from db.database import DatabaseManager, to_minor
from migration.migrar_compras_cuotas import (
    COL_BANCO,
    COL_CONCEPTO,
    COL_CUOTAS_PAGADAS,
    COL_CUOTAS_TOTALES,
    COL_FECHA_COMPRA,
    COL_FECHA_PRIMER_PAGO,
    COL_MONEDA,
    COL_PRECIO_TOTAL,
    COL_SUBCATEGORIA,
    COLUMNAS_OBLIGATORIAS,
    COLUMNAS_OPCIONALES,
    FILAS_BUSQUEDA_HEADER,
    HOJA,
    MARCAR_CUOTAS_YA_PAGADAS,
    MONEDA_DEFAULT,
    PROYECTAR_DESDE_PRIMER_PAGO,
    MigracionError,
    _asegurar_cuenta,
    _categorias_activas,
    _clave,
    _mapa_categorias,
    _parsear_entero,
    _parsear_fecha,
    _parsear_monto,
    _sumar_meses,
    _texto,
)

RAIZ = Path(__file__).resolve().parent.parent
DB_PATH = RAIZ / "data" / "deltabalanceNR.db"
# La DB real de la app y la de Bruno: este script nunca las toca.
DBS_PROHIBIDAS = (RAIZ / "data" / "deltabalance.db", RAIZ / "data" / "deltabalanceBZ.db")
SOURCES_DIR = Path(__file__).resolve().parent / "sources"
XLSX_PATH = SOURCES_DIR / "Finanzas.xlsx"
NOTA_MIGRACION = "MIGRADO DESDE EXCEL NOE — COMPRAS EN CUOTAS, FILA {fila}"

# --- Cuentas ---
CUENTAS_DEBITO = ["BERSA"]
# Tarjeta de crédito → cuenta débito desde la que se paga (None = ninguna).
CUENTAS_CREDITO: dict[str, Optional[str]] = {
    "BERSA VISA EXT RP": "BERSA",  # extensión tarjeta madre
    "BZ":                None,     # compras de Bruno — sin cuenta débito asociada
}

# --- Mapeos (valor de la planilla → nombre en la DB) ---
MAPA_BANCO = {
    "VISA": "BERSA VISA EXT RP",
    "OTRO": "BZ",
}
MAPA_CATEGORIA = {
    "TARJETA": "EGRESO VARIABLE",
}
# Cualquier subcategoría fuera de MAPA_CATEGORIA (sin aviso: es lo esperado).
CATEGORIA_FALLBACK = "EGRESO VARIABLE"

# Nombres aceptados para la cantidad total de cuotas, en orden de
# preferencia. Por nombre EXACTO, no por "contiene": "CANTIDAD DE CUOTAS
# PAGADAS" también contiene "CANTIDAD DE CUOTAS".
ALIAS_CUOTAS_TOTALES = (COL_CUOTAS_TOTALES, "CANTIDAD DE CUOTAS")

# --- Categorías a crear si no existen (principal y tipo del script de Bruno) ---
CATEGORIAS_NUEVAS = [
    ("EGRESOS VARIABLES", "EGRESO VARIABLE", "egreso"),
]


# ============================================================
# CHEQUEOS DE CONFIGURACIÓN (antes de tocar la DB)
# ============================================================

def _validar_configuracion() -> None:
    """Cada tarjeta de MAPA_BANCO está en CUENTAS_CREDITO, y cada débito de CUENTAS_CREDITO en CUENTAS_DEBITO."""
    sueltas = sorted(
        {f"{t} (tarjeta)" for t in MAPA_BANCO.values() if t not in CUENTAS_CREDITO}
        | {f"{d} (débito)" for d in CUENTAS_CREDITO.values() if d is not None and d not in CUENTAS_DEBITO}
    )
    if sueltas:
        raise MigracionError(f"Cuentas usadas que no están configuradas: {', '.join(sueltas)}.")


# ============================================================
# CUENTAS Y CATEGORÍA
# ============================================================

def _crear_cuentas(conn: sqlite3.Connection, moneda_ars_id: int) -> dict[str, str]:
    ids: dict[str, str] = {}
    print("\n[CUENTAS DÉBITO]")
    for nombre in CUENTAS_DEBITO:
        ids[nombre] = _asegurar_cuenta(conn, nombre, "debito", moneda_ars_id)
    print("\n[CUENTAS CRÉDITO]")
    for nombre, debito in CUENTAS_CREDITO.items():
        ids[nombre] = _asegurar_cuenta(
            conn, nombre, "credito", moneda_ars_id,
            cuenta_pago_id=ids[debito] if debito is not None else None, nombre_debito=debito,
        )
    return ids


def _crear_categorias(conn: sqlite3.Connection) -> None:
    print("\n[CATEGORÍAS]")
    for principal, sub, tipo in CATEGORIAS_NUEVAS:
        existente = next(
            (c for c in _categorias_activas(conn) if _clave(c["subcategoria"]) == _clave(sub)), None,
        )
        if existente is not None:
            ubicacion = f"{existente['categoria_principal']} · {existente['subcategoria']}"
            como = "" if ubicacion == f"{principal} · {sub}" else f" (como '{ubicacion}' — se reusa, no se duplica)"
            print(f"  ⏭️  {sub} ya existe{como}")
            continue
        cur = conn.execute(
            "INSERT OR IGNORE INTO categorias (categoria_principal, subcategoria, tipo) VALUES (?, ?, ?);",
            (principal, sub, tipo),
        )
        if cur.rowcount == 0:
            # Existe con ese nombre exacto pero desactivada: no se revierte un
            # soft-delete hecho a propósito — las filas quedan como error.
            print(f"  ⚠️  {principal} · {sub} existe DESACTIVADA — no se reactiva")
        else:
            print(f"  ✅ {sub} creada")


# ============================================================
# FILAS
# ============================================================

def _encontrar_header(ws) -> tuple[int, dict[str, int]]:
    """
    Como migrar_compras_cuotas._encontrar_header(), pero la cantidad de
    cuotas se busca por ALIAS_CUOTAS_TOTALES: la que aparezca queda en
    columnas[COL_CUOTAS_TOTALES] (lo que lee _procesar_filas()).
    """
    for numero, valores in enumerate(
        ws.iter_rows(min_row=1, max_row=FILAS_BUSQUEDA_HEADER, values_only=True), start=1,
    ):
        textos = [_texto(v) for v in valores]
        if COL_CONCEPTO in textos and COL_BANCO in textos:
            columnas: dict[str, int] = {}
            for indice, texto in enumerate(textos):
                if texto:
                    columnas.setdefault(texto, indice)
            alias = next((a for a in ALIAS_CUOTAS_TOTALES if a in columnas), None)
            if alias is not None:
                columnas.setdefault(COL_CUOTAS_TOTALES, columnas[alias])
            faltantes = [
                " o ".join(ALIAS_CUOTAS_TOTALES) if c == COL_CUOTAS_TOTALES else c
                for c in COLUMNAS_OBLIGATORIAS if c not in columnas
            ]
            if faltantes:
                raise MigracionError(f"A la hoja '{HOJA}' le faltan columnas: {', '.join(faltantes)}.")
            return numero, columnas
    raise MigracionError(
        f"No se encontró la fila de encabezados (con {COL_CONCEPTO} y {COL_BANCO}) "
        f"en las primeras {FILAS_BUSQUEDA_HEADER} filas de '{HOJA}'."
    )


def _procesar_filas(
    conn: sqlite3.Connection,
    ids_cuentas: dict[str, str],
    categorias: dict[str, sqlite3.Row],
    monedas: dict[str, sqlite3.Row],
) -> dict[str, int]:
    """Mismo recorrido que migrar_compras_cuotas._procesar_filas(), con los mapeos de Noe (sin inferencia de categoría)."""
    cont = {
        "total": 0, "importadas": 0, "importadas_todas_pagadas": 0, "ya_existian": 0,
        "saltadas_banco": 0, "errores": 0, "fallback": 0,
    }
    mapa_banco = {_clave(k): v for k, v in MAPA_BANCO.items()}
    mapa_categoria = {_clave(k): v for k, v in MAPA_CATEGORIA.items()}

    print(f"\n[FILAS] Procesando hoja {HOJA}...")
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    try:
        if HOJA not in wb.sheetnames:
            raise MigracionError(f"La planilla no tiene la hoja '{HOJA}'.")
        ws = wb[HOJA]
        fila_header, columnas = _encontrar_header(ws)
        for col in COLUMNAS_OPCIONALES:
            if col not in columnas:
                print(f"  ⚠️  La hoja no tiene la columna '{col}' — se sigue sin ella (ver docstring de migrar_compras_cuotas.py).")

        for fila_excel, valores in enumerate(
            ws.iter_rows(min_row=fila_header + 1, values_only=True), start=fila_header + 1,
        ):
            def celda(col: str) -> Any:
                indice = columnas.get(col)
                return valores[indice] if indice is not None and indice < len(valores) else None

            concepto = _texto(celda(COL_CONCEPTO))
            if not concepto:
                continue  # fila vacía
            cont["total"] += 1
            etiqueta = f"[{fila_excel:03d}] {concepto}"

            def error(motivo: str) -> None:
                cont["errores"] += 1
                print(f"  ❌ {etiqueta} | {motivo} — salteando")

            cuotas_totales = _parsear_entero(celda(COL_CUOTAS_TOTALES))
            cuotas_pagadas = _parsear_entero(celda(COL_CUOTAS_PAGADAS)) or 0

            # --- Banco (OTRO = la tarjeta de Bruno, no se saltea) ---
            banco = _texto(celda(COL_BANCO))
            nombre_tarjeta = mapa_banco.get(_clave(banco))
            if nombre_tarjeta is None:
                cont["saltadas_banco"] += 1
                print(f'  ⚠️  {etiqueta} | BANCO DESCONOCIDO: "{banco}" — salteando')
                continue
            cuenta_id = ids_cuentas[nombre_tarjeta]

            # --- Monto, fecha, cuotas ---
            precio = _parsear_monto(celda(COL_PRECIO_TOTAL))
            if precio is None:
                error(f"PRECIO TOTAL INVÁLIDO: {celda(COL_PRECIO_TOTAL)!r}")
                continue

            fecha_compra = _parsear_fecha(celda(COL_FECHA_COMPRA))
            if fecha_compra is None:
                error(f"FECHA COMPRA INVÁLIDA: {celda(COL_FECHA_COMPRA)!r}")
                continue

            if cuotas_totales is None:
                error(f"CANTIDAD DE CUOTAS TOTALES INVÁLIDA: {celda(COL_CUOTAS_TOTALES)!r}")
                continue
            avisos: list[str] = []
            if cuotas_totales < 1:
                avisos.append(f"CANTIDAD DE CUOTAS {cuotas_totales} → 1")
                cuotas_totales = 1
            cuotas_pagadas = max(0, min(cuotas_pagadas, cuotas_totales))

            # --- Categoría: mapa → fallback (sin aviso, ver docstring) ---
            destino = mapa_categoria.get(_clave(_texto(celda(COL_SUBCATEGORIA))))
            if destino is None:
                cont["fallback"] += 1
                destino = CATEGORIA_FALLBACK
            categoria = categorias.get(_clave(destino))
            if categoria is None:
                error(f"LA CATEGORÍA DESTINO '{destino}' NO EXISTE (ACTIVA) EN LA DB")
                continue

            # --- Moneda ---
            moneda_txt = _texto(celda(COL_MONEDA))
            moneda = monedas.get(_clave(moneda_txt))
            if moneda is None:
                moneda = monedas[_clave(MONEDA_DEFAULT)]
                avisos.append(f"MONEDA DESCONOCIDA: \"{moneda_txt or '(vacía)'}\" → {MONEDA_DEFAULT}")

            monto_total_minor = to_minor(precio, moneda["decimales"])
            monto_por_cuota_minor = round(monto_total_minor / cuotas_totales)

            # --- Idempotencia ---
            existente = conn.execute(
                """
                SELECT id FROM compras_cuotas
                WHERE concepto = ? AND fecha_compra = ? AND cuenta_id = ? AND monto_total_minor = ?
                LIMIT 1;
                """,
                (concepto, fecha_compra, cuenta_id, monto_total_minor),
            ).fetchone()
            if existente is not None:
                cont["ya_existian"] += 1
                print(f"  ⏭️  {etiqueta} | YA EXISTE (compra #{existente['id']}) — salteando")
                continue

            # --- Cronograma ---
            base_cronograma = fecha_compra
            if PROYECTAR_DESDE_PRIMER_PAGO:
                primer_pago = _parsear_fecha(celda(COL_FECHA_PRIMER_PAGO))
                if primer_pago is not None:
                    base_cronograma = primer_pago
                elif COL_FECHA_PRIMER_PAGO in columnas:
                    avisos.append("SIN FECHA PRIMER PAGO VÁLIDA — cuotas proyectadas desde FECHA COMPRA")
            anio_base, mes_base = int(base_cronograma[:4]), int(base_cronograma[5:7])

            # --- La tarjeta tiene que operar en la moneda de la compra ---
            cur = conn.execute(
                "INSERT OR IGNORE INTO cuentas_saldos (cuenta_id, moneda_id, saldo_inicial_minor) VALUES (?, ?, 0);",
                (cuenta_id, moneda["id"]),
            )
            if cur.rowcount:
                avisos.append(f"{nombre_tarjeta} AHORA OPERA TAMBIÉN EN {moneda['codigo']} (cuentas_saldos)")

            # --- Insert: compra + cuotas ---
            cur = conn.execute(
                """
                INSERT INTO compras_cuotas
                    (fecha_compra, concepto, cuenta_id, categoria_id, moneda_id,
                     monto_total_minor, total_cuotas, monto_por_cuota_minor, notas)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
                """,
                (
                    fecha_compra, concepto, cuenta_id, categoria["id"], moneda["id"],
                    monto_total_minor, cuotas_totales, monto_por_cuota_minor,
                    NOTA_MIGRACION.format(fila=fila_excel),
                ),
            )
            # El id es un UUID (DEFAULT de la columna): lastrowid es el rowid interno, no el id.
            compra_id = conn.execute("SELECT id FROM compras_cuotas WHERE rowid = ?;", (cur.lastrowid,)).fetchone()[0]
            for numero in range(1, cuotas_totales + 1):
                anio, mes = _sumar_meses(anio_base, mes_base, numero - 1)
                estado = "pagado" if (MARCAR_CUOTAS_YA_PAGADAS and numero <= cuotas_pagadas) else "pendiente"
                conn.execute(
                    """
                    INSERT INTO cuotas_credito
                        (compra_id, numero_cuota, mes_proyectado, anio_proyectado, monto_cuota_minor, estado)
                    VALUES (?, ?, ?, ?, ?, ?);
                    """,
                    (compra_id, numero, mes, anio, monto_por_cuota_minor, estado),
                )

            cont["importadas"] += 1
            todas_pagadas = MARCAR_CUOTAS_YA_PAGADAS and cuotas_pagadas == cuotas_totales
            if todas_pagadas:
                cont["importadas_todas_pagadas"] += 1
            for aviso in avisos:
                print(f"  ⚠️  {etiqueta} | {aviso}")
            if todas_pagadas:
                ya_pagadas = " (TODAS PAGADAS)"
            elif MARCAR_CUOTAS_YA_PAGADAS and cuotas_pagadas:
                ya_pagadas = f" ({cuotas_pagadas} YA PAGADAS)"
            else:
                ya_pagadas = ""
            print(
                f"  ✅ {etiqueta} | {nombre_tarjeta} | {moneda['codigo']} ${precio:,.2f} | "
                f"{cuotas_totales} CUOTAS{ya_pagadas} | {fecha_compra} | 1ª CUOTA {mes_base:02d}/{anio_base}"
            )
    finally:
        wb.close()
    return cont


# ============================================================
# MAIN
# ============================================================

def _mostrar_ruta(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(RAIZ))
    except ValueError:
        return str(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Migra las compras en cuotas desde la planilla de Noe.")
    parser.add_argument("--db-path", default=None, help=f"Ruta a la DB (default: {_mostrar_ruta(DB_PATH)}).")
    parser.add_argument("--confirmar", action="store_true", help="Aplica los cambios. Sin esto, solo dry-run.")
    args = parser.parse_args()

    db_path = Path(args.db_path) if args.db_path else DB_PATH
    for prohibida in DBS_PROHIBIDAS:
        if db_path.resolve() == prohibida.resolve():
            sys.exit(f"❌ Este script nunca corre contra {_mostrar_ruta(prohibida)} (es de Noe: {_mostrar_ruta(DB_PATH)}).")

    if not XLSX_PATH.exists():
        sys.exit(f"❌ No se encontró la planilla de Noe: {_mostrar_ruta(XLSX_PATH)}.")

    try:
        _validar_configuracion()
    except MigracionError as err:
        sys.exit(f"❌ {err}")

    modo = "APLICANDO CAMBIOS" if args.confirmar else "DRY-RUN (nada se escribe)"
    print(f"[INICIO] Migración de compras en cuotas de NOE — {modo}")
    print(f"[XLSX] {_mostrar_ruta(XLSX_PATH)}")
    print(f"[DB] Conectando a {_mostrar_ruta(db_path)}")

    db = DatabaseManager(db_path)
    db.inicializar()  # idempotente: schema + migraciones; seed solo si la DB es nueva
    conn = db.conn
    # inicializar() puede dejar abierta una transacción implícita: se cierra
    # para que el rollback del dry-run o de un error deshaga SOLO lo de este script.
    conn.commit()

    if args.confirmar:
        # El backup copia solo el archivo .db — con WAL, lo que siga en el -wal no estaría sin este checkpoint.
        conn.execute("PRAGMA wal_checkpoint(TRUNCATE);")
        db.hacer_backup()

    try:
        monedas = {_clave(m["codigo"]): m for m in conn.execute("SELECT id, codigo, decimales FROM monedas;")}
        moneda_ars = monedas.get(_clave(MONEDA_DEFAULT))
        if moneda_ars is None:
            raise MigracionError(f"La moneda default {MONEDA_DEFAULT} no existe en la DB.")

        ids_cuentas = _crear_cuentas(conn, moneda_ars["id"])
        _crear_categorias(conn)
        categorias = _mapa_categorias(conn)

        destinos = set(MAPA_CATEGORIA.values()) | {CATEGORIA_FALLBACK}
        faltantes = sorted(d for d in destinos if _clave(d) not in categorias)
        if faltantes:
            print(f"\n  ⚠️  Categorías destino que NO existen activas en la DB (sus filas van a Errores): {', '.join(faltantes)}")

        cont = _procesar_filas(conn, ids_cuentas, categorias, monedas)
    except Exception as err:
        conn.rollback()
        db.desconectar()
        print(f"\n❌ ERROR — rollback completo, no se escribió nada: {err}")
        if isinstance(err, MigracionError):
            sys.exit(1)
        raise

    if args.confirmar:
        conn.commit()
    else:
        conn.rollback()
    db.desconectar()

    print(f"\n[FIN] NOE — {modo}")
    print(f"  Total filas procesadas : {cont['total']}")
    print(f"  Importadas             : {cont['importadas']}")
    print(f"    con todas las cuotas pagadas (históricas) : {cont['importadas_todas_pagadas']}")
    print(f"  Ya existían            : {cont['ya_existian']}")
    print(f"  Saltadas (banco desc.) : {cont['saltadas_banco']}")
    print(f"  Errores                : {cont['errores']}")
    print(f"  Categoría por fallback : {cont['fallback']} ({CATEGORIA_FALLBACK}, esperado)")
    if not args.confirmar:
        print("\nNada se escribió — revisá el reporte y corré de nuevo con --confirmar para aplicar.")


if __name__ == "__main__":
    main()
